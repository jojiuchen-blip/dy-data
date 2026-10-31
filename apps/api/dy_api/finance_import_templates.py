"""Chinese finance templates and their backwards-compatible input contract."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


FINANCE_IMPORT_TEMPLATE_FIELDS = {
    "BASIC_INFO": ["storeId", "storeName", "sapCode", "importedAt"],
    "PROMOTION_FACTORY_RESULT": [
        "invoiceNumber", "reviewResult", "rejectionReason", "settlementDate", "settlementAmountCent",
    ],
    "MANAGEMENT_FACTORY_RESULT": [
        "storeId", "statementMonth", "storeName", "invoiceNumber", "invoiceDate", "deductionDate", "deductionAmountCent",
    ],
    "SAP_CONFIRMATION": [
        "storeId", "storeName", "financeInitialSap", "serviceStoreCode", "finalSapCode", "factoryConfirmationResult", "confirmedAt",
    ],
}
FINANCE_IMPORT_FIELD_LABELS = {
    "businessKey": "业务唯一键", "file": "文件",
    "storeId": "门店ID", "storeName": "门店名称", "sapCode": "SAP编码", "importedAt": "资料导入时间（含时区）",
    "invoiceNumber": "发票号码", "reviewResult": "厂家审核结果", "rejectionReason": "审核不通过原因",
    "settlementDate": "结算日期", "settlementAmountCent": "结算金额（分）", "statementMonth": "业务账期（YYYY-MM）",
    "invoiceDate": "开票日期", "deductionDate": "厂家扣款日期", "deductionAmountCent": "厂家扣款金额（分）",
    "financeInitialSap": "财务初始SAP编码", "serviceStoreCode": "服务门店编码", "finalSapCode": "最终有效SAP编码",
    "factoryConfirmationResult": "厂家确认结果", "confirmedAt": "厂家确认时间（含时区）",
}
FINANCE_IMPORT_TYPE_LABELS = {
    "BASIC_INFO": "门店基础信息", "PROMOTION_FACTORY_RESULT": "推广服务费厂家结果",
    "MANAGEMENT_FACTORY_RESULT": "管理服务费厂家结果", "SAP_CONFIRMATION": "SAP确认",
}
FINANCE_IMPORT_VALUE_ALIASES = {
    "reviewResult": {"审核通过并已结算": "APPROVED_SETTLED", "审核不通过需重新上传": "REJECTED_REUPLOAD"},
    "factoryConfirmationResult": {"已确认": "CONFIRMED", "确认不通过": "REJECTED"},
}
FINANCE_IMPORT_LEGACY_IGNORED_FIELDS = {
    "BASIC_INFO": {"sapStatus", "sap_status", "sapConfirmationStatus", "sapCodeConfirmationStatus", "factoryConfirmed", "factoryConfirmationResult", "SAP确认状态", "SAP编码确认状态"},
}
FINANCE_IMPORT_EXAMPLE_MARKER = "示例标记（请勿导入）"
FINANCE_IMPORT_IDENTIFIER_FIELDS = {
    "storeId", "sapCode", "invoiceNumber", "financeInitialSap", "serviceStoreCode", "finalSapCode",
}
_EXAMPLES = {
    "BASIC_INFO": ["00000000000000000001", "虚构示例服务店（非真实门店）", "0000123456", "2026-09-01T10:00:00+08:00"],
    "PROMOTION_FACTORY_RESULT": ["00000000000000000001", "审核通过并已结算", "", "2026-10-08", "12345"],
    "MANAGEMENT_FACTORY_RESULT": ["00000000000000000002", "2026-09", "虚构示例管理服务店（非真实门店）", "00000000000000000002", "2026-10-08", "2026-10-09", "6789"],
    "SAP_CONFIRMATION": ["00000000000000000003", "虚构示例SAP服务店（非真实门店）", "0000123000", "0000004567", "0000123999", "已确认", "2026-10-09T11:30:00+08:00"],
}
_NOTES = {
    "storeId": "必填；系统门店ID，按文本填写；不得用SAP编码或门店名称替代。示例ID不存在。",
    "storeName": "必填；与门店ID对应的系统当前名称精确一致，仅作身份复核。",
    "sapCode": "首次资料必填；1—128字符，按文本填写并保留前导零。已有资料时空白或未提供保留当前SAP，不清空、不确认新编码。",
    "importedAt": "必填；ISO 8601时间且含时区，例如2026-09-01T10:00:00+08:00。",
    "invoiceNumber": "必填；20位数字，必须按文本填写并保留前导零；只精确匹配当前业务记录。",
    "reviewResult": "必填；审核通过并已结算 / 审核不通过需重新上传；兼容APPROVED_SETTLED / REJECTED_REUPLOAD。",
    "rejectionReason": "审核不通过时必填；审核通过可空，新结果不带退回原因，旧事件保留；不修改原发票金额。",
    "settlementDate": "审核通过时必填；YYYY-MM-DD。审核不通过可空。",
    "settlementAmountCent": "审核通过时必填；正整数，单位为分（12345分=123.45元），须等于原发票全额；不支持部分结算。",
    "statementMonth": "必填；YYYY-MM，须与页面所选导入批次账期一致，填写厂家结果对应的账期。",
    "invoiceDate": "必填；YYYY-MM-DD。",
    "deductionDate": "必填；YYYY-MM-DD，填写厂家实际扣款日期。",
    "deductionAmountCent": "必填；正整数，单位为分（6789分=67.89元），须等于有效确认金额；不支持部分扣款。",
    "financeInitialSap": "必填；财务初始SAP编码，按文本填写，最长128字符。",
    "serviceStoreCode": "必填；服务门店编码，按文本填写，最长128字符。",
    "finalSapCode": "必填；厂家确认后的最终有效SAP编码，按文本填写，最长128字符。",
    "factoryConfirmationResult": "必填；已确认 / 确认不通过；兼容CONFIRMED / REJECTED。仅本独立SAP确认模板可提交该结果。",
    "confirmedAt": "必填；含时区的ISO 8601时间，填写实际厂家确认时间。",
}


def finance_import_headers(import_type: str) -> list[str]:
    """Return the Chinese headers in the formal field order."""
    return [FINANCE_IMPORT_FIELD_LABELS[field] for field in FINANCE_IMPORT_TEMPLATE_FIELDS[import_type]]


def finance_import_header_keys(import_type: str, headers: list[str]) -> list[str | None]:
    """Resolve Chinese/English headers, rejecting ambiguous or unknown columns."""
    fields = FINANCE_IMPORT_TEMPLATE_FIELDS[import_type]
    aliases = {field: field for field in fields}
    aliases.update({FINANCE_IMPORT_FIELD_LABELS[field]: field for field in fields})
    # Earlier Chinese labels are accepted as explicit aliases, never positionally.
    aliases.update({label: field for label, field in {
        "门店编号": "storeId", "SAP 编码": "sapCode", "导入时间": "importedAt",
        "业务账期": "statementMonth", "结算金额（单位：分）": "settlementAmountCent",
        "扣款金额（分）": "deductionAmountCent", "确认时间": "confirmedAt",
    }.items() if field in fields})
    ignored = FINANCE_IMPORT_LEGACY_IGNORED_FIELDS.get(import_type, set())
    keys: list[str | None] = []
    seen: set[str] = set()
    for header in headers:
        if header == FINANCE_IMPORT_EXAMPLE_MARKER:
            raise ValueError("填写示例不能导入；请在“数据导入”表填写真实资料")
        if header in ignored:
            keys.append(None)
            continue
        key = aliases.get(header)
        if key is None:
            raise ValueError(f"无法识别表头“{header or '空表头'}”，请使用对应类型的标准模板")
        if key in seen:
            raise ValueError(f"字段“{FINANCE_IMPORT_FIELD_LABELS[key]}”重复，中文和英文别名不能同时出现")
        seen.add(key)
        keys.append(key)
    return keys


def build_finance_import_workbook(import_type: str) -> bytes:
    """Build a blank import sheet plus isolated, explicitly fictional guidance."""
    workbook = Workbook()
    data = workbook.active
    data.title = "数据导入"
    headers = finance_import_headers(import_type)
    data.append(headers)
    example = workbook.create_sheet("填写示例")
    example.append([*headers, FINANCE_IMPORT_EXAMPLE_MARKER])
    example.append([*_EXAMPLES[import_type], "全部为虚构示例，请勿提交"])
    notes = workbook.create_sheet("填写说明")
    notes.append(["填写说明（不参与导入）", FINANCE_IMPORT_TYPE_LABELS[import_type]])
    notes.append(["使用方法", "只在“数据导入”表填写；“填写示例”和本说明表不会导入。样例不代表真实业务资料。"])
    notes.append(["匹配与空白", "按稳定ID精确匹配；空白语义以各字段说明为准；基础信息SAP空白保留已有值。任一错误整批不写入。"])
    notes.append(["旧模板兼容", "旧英文CSV/XLSX按字段名兼容；基础信息旧SAP确认状态列忽略，不改变已有SAP确认事实。"])
    notes.append(["正式生效", "上传仅预览；核对全部原值、新值及允许变更后提交。提交后刷新正式记录；不改变历史账单/发票快照。"])
    for field in FINANCE_IMPORT_TEMPLATE_FIELDS[import_type]:
        notes.append([FINANCE_IMPORT_FIELD_LABELS[field], _NOTES[field]])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
        for row in sheet:
            for cell in row:
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            dimension = sheet.column_dimensions[column[0].column_letter]
            dimension.width = 28 if sheet != notes else (30 if column[0].column == 1 else 100)
            dimension.number_format = "@"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
