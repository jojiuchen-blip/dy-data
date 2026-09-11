"""Parse approved ranking rosters and publish complete sidecar versions."""

from __future__ import annotations

import csv
from datetime import date, datetime, timezone
import math
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from apps.api.dy_api.models import DimStorePoiMapping
from apps.api.dy_api.ranking_configuration import (
    DATA_START,
    RANKING_WRITE_LOCK,
    publish_configuration,
)
from apps.api.dy_api.ranking_schema_v1 import eligibility, org_history


MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_STORE_COUNT = 20_000
MISSING_ORGANIZATION_NAME = "待补充归属"
ORGANIZATION_COLUMN_ALIASES = {
    "poi_id": ("所属账户关联poi_ID",),
    "service_store_code": ("服务店编码",),
    "store_name": ("服务店名称", "门店名称"),
    "group_name": ("所属集团", "集团"),
    "service_center_name": ("服务中心", "所属服务中心"),
    "district_name": ("大区", "所属大区"),
    "area_name": ("区域", "所属区域"),
}


def _utc(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            raise ValueError("名单中的编码和文本单元格不得使用非整数数值")
        return str(int(value))
    return str(value).strip()


def _validate_file(path: Path) -> str:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise ValueError(f"名单文件无法读取: {path.name}") from exc
    if not path.is_file():
        raise ValueError(f"名单文件不存在: {path.name}")
    if size > MAX_IMPORT_BYTES:
        raise ValueError("名单文件不能超过10 MiB")
    extension = path.suffix.lower()
    if extension not in {".csv", ".xlsx", ".xlsm"}:
        raise ValueError("名单文件仅支持XLSX、XLSM或CSV格式")
    if extension in {".xlsx", ".xlsm"}:
        try:
            with ZipFile(path) as archive:
                expanded_size = sum(member.file_size for member in archive.infolist())
        except (BadZipFile, OSError) as exc:
            raise ValueError("XLSX文件无法解析") from exc
        if expanded_size > MAX_XLSX_UNCOMPRESSED_BYTES:
            raise ValueError("XLSX文件解压后不能超过100 MiB")
    return extension


def _validate_headers(headers: tuple[str, ...]) -> None:
    if not headers or not all(headers):
        raise ValueError("名单文件存在空列名")
    if len(set(headers)) != len(headers):
        raise ValueError("名单文件存在重复列名")


def _has_cell_value(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _read_csv_rows(path: Path) -> list[dict[str, object]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            reader = csv.reader(source)
            raw_header = next(reader, None)
            if raw_header is None:
                raise ValueError("名单文件为空")
            headers = tuple(_cell_text(value) for value in raw_header)
            _validate_headers(headers)
            result: list[dict[str, object]] = []
            for row_number, source_row in enumerate(reader, start=2):
                values = tuple(source_row)
                if not any(_has_cell_value(value) for value in values):
                    continue
                if len(values) != len(headers):
                    raise ValueError(f"名单第{row_number}行列数不正确")
                result.append(dict(zip(headers, values)))
                if len(result) > MAX_STORE_COUNT:
                    raise ValueError("名单最多包含20000家门店")
            return result
    except UnicodeDecodeError as exc:
        raise ValueError("CSV文件必须使用UTF-8编码") from exc
    except csv.Error as exc:
        raise ValueError("CSV文件无法解析") from exc


def _read_xlsx_rows(path: Path) -> list[dict[str, object]]:
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:  # noqa: BLE001 - malformed workbook boundary.
        raise ValueError("XLSX文件无法解析") from exc
    try:
        iterator = workbook.active.iter_rows(values_only=True)
        raw_header = next(iterator, None)
        if raw_header is None:
            raise ValueError("名单文件为空")
        headers = tuple(_cell_text(value) for value in raw_header)
        _validate_headers(headers)
        result: list[dict[str, object]] = []
        for row_number, source_row in enumerate(iterator, start=2):
            values = tuple(source_row)
            if not any(_has_cell_value(value) for value in values):
                continue
            if len(values) != len(headers):
                raise ValueError(f"名单第{row_number}行列数不正确")
            result.append(dict(zip(headers, values)))
            if len(result) > MAX_STORE_COUNT:
                raise ValueError("名单最多包含20000家门店")
        return result
    finally:
        workbook.close()


def _read_rows(path: Path) -> list[dict[str, object]]:
    extension = _validate_file(path)
    rows = (
        _read_csv_rows(path)
        if extension == ".csv"
        else _read_xlsx_rows(path)
    )
    if not rows:
        raise ValueError("名单至少包含一家门店")
    return rows


def _resolve_store_ids(session: Session, poi_ids: set[str]) -> dict[str, str]:
    candidates: dict[str, set[str]] = {poi_id: set() for poi_id in poi_ids}
    ordered_pois = sorted(poi_ids)
    for start in range(0, len(ordered_pois), 500):
        batch = ordered_pois[start : start + 500]
        rows = session.execute(
            select(DimStorePoiMapping.poi_id, DimStorePoiMapping.store_id).where(
                DimStorePoiMapping.poi_id.in_(batch)
            )
        )
        for poi_id, store_id in rows:
            normalized_poi = _cell_text(poi_id)
            normalized_store = _cell_text(store_id)
            if normalized_poi in candidates and normalized_store:
                candidates[normalized_poi].add(normalized_store)
    resolved: dict[str, str] = {}
    for poi_id in ordered_pois:
        stores = candidates[poi_id]
        if not stores:
            raise ValueError(f"POI未匹配DimStore: {poi_id}")
        if len(stores) != 1:
            raise ValueError(f"POI对应多个门店，无法确定唯一DimStore: {poi_id}")
        resolved[poi_id] = next(iter(stores))
    return resolved


def _required(row: dict[str, object], column: str, row_number: int) -> str:
    value = _cell_text(row[column])
    if not value:
        raise ValueError(f"名单第{row_number}行的{column}不能为空")
    return value


def _required_identifier(row: dict[str, object], column: str, row_number: int) -> str:
    source = row[column]
    if isinstance(source, bool) or not isinstance(source, (str, int)):
        raise ValueError(f"名单第{row_number}行的{column}必须为精确字符串或整数")
    value = str(source).strip()
    if not value:
        raise ValueError(f"名单第{row_number}行的{column}不能为空")
    return value


def _organization_value(row: dict[str, object], column: str) -> str:
    return _cell_text(row[column]) or MISSING_ORGANIZATION_NAME


def _resolve_columns(
    headers: set[str], aliases_by_field: dict[str, tuple[str, ...]]
) -> dict[str, str]:
    resolved: dict[str, str] = {}
    for field, aliases in aliases_by_field.items():
        matches = [alias for alias in aliases if alias in headers]
        if not matches:
            raise ValueError(f"名单缺少必需列: {'/'.join(aliases)}")
        if len(matches) > 1 and field != "store_name":
            raise ValueError(f"名单包含重复业务列: {'/'.join(matches)}")
        resolved[field] = matches[0]
    return resolved


def _parse_organizations(session: Session, path: Path) -> list[dict[str, str]]:
    rows = _read_rows(path)
    columns = _resolve_columns(set(rows[0]), ORGANIZATION_COLUMN_ALIASES)
    poi_ids: list[str] = []
    seen_pois: set[str] = set()
    for row_number, row in enumerate(rows, start=2):
        poi_id = _required_identifier(row, columns["poi_id"], row_number)
        _required_identifier(row, columns["service_store_code"], row_number)
        if poi_id in seen_pois:
            raise ValueError(f"组织名单存在重复POI: {poi_id}")
        poi_ids.append(poi_id)
        seen_pois.add(poi_id)
    store_ids = _resolve_store_ids(session, set(poi_ids))
    organizations: list[dict[str, str]] = []
    for row_number, row in enumerate(rows, start=2):
        poi_id = _required_identifier(row, columns["poi_id"], row_number)
        group_name = _organization_value(row, columns["group_name"])
        group_key = (
            ""
            if group_name == MISSING_ORGANIZATION_NAME
            else _cell_text(row.get("集团编码", ""))
        )
        organizations.append(
            {
                "store_id": store_ids[poi_id],
                "service_store_code": _required_identifier(
                    row, columns["service_store_code"], row_number
                ),
                "store_name": _required(row, columns["store_name"], row_number),
                "group_name": group_name,
                "group_key": group_key,
                "service_center_name": _required(
                    row, columns["service_center_name"], row_number
                ),
                "district_name": _organization_value(row, columns["district_name"]),
                "area_name": _organization_value(row, columns["area_name"]),
            }
        )
    return organizations


def _parse_eligibility_pois(path: Path) -> list[str]:
    rows = _read_rows(path)
    if "所属账户关联poi_ID" not in rows[0]:
        raise ValueError("资格名单缺少必需列: 所属账户关联poi_ID")
    result: list[str] = []
    seen: set[str] = set()
    for row_number, row in enumerate(rows, start=2):
        poi_id = _required_identifier(row, "所属账户关联poi_ID", row_number)
        if poi_id in seen:
            raise ValueError(f"资格名单存在重复POI: {poi_id}")
        seen.add(poi_id)
        result.append(poi_id)
    return result


def _latest_organizations(session: Session) -> tuple[list[dict[str, str]], str | None, datetime | None]:
    latest = session.execute(
        select(org_history.c.mapping_version, org_history.c.effective_from)
        .order_by(org_history.c.effective_from.desc(), org_history.c.mapping_version.desc())
        .limit(1)
    ).first()
    if latest is None:
        return [], None, None
    mapping_version, effective_from = latest
    rows = session.execute(
        select(org_history).where(org_history.c.mapping_version == mapping_version)
    ).mappings()
    organizations = [
        {
            "store_id": row["store_id"],
            "service_store_code": row["service_store_code"],
            "store_name": row["store_name"],
            "group_name": row["group_name"],
            "group_key": row["group_key"],
            "service_center_name": row["service_center_name"],
            "district_name": row["district_name"],
            "area_name": row["area_name"],
        }
        for row in rows
    ]
    return organizations, mapping_version, _utc(effective_from)


def _latest_eligibility(session: Session) -> tuple[list[str], str | None, datetime | None]:
    latest = session.execute(
        select(eligibility.c.eligibility_version, eligibility.c.effective_from)
        .where(eligibility.c.product_scope == "精诚养车")
        .order_by(eligibility.c.effective_from.desc(), eligibility.c.eligibility_version.desc())
        .limit(1)
    ).first()
    if latest is None:
        return [], None, None
    eligibility_version, effective_from = latest
    codes = list(
        session.scalars(
            select(eligibility.c.service_store_code).where(
                eligibility.c.eligibility_version == eligibility_version,
                eligibility.c.product_scope == "精诚养车",
            )
        )
    )
    return codes, eligibility_version, _utc(effective_from)


def _eligible_codes(
    session: Session,
    poi_ids: list[str],
    organizations: list[dict[str, str]],
) -> list[str]:
    resolved = _resolve_store_ids(session, set(poi_ids))
    code_by_store: dict[str, str] = {}
    for row in organizations:
        store_id = row["store_id"]
        if store_id in code_by_store and code_by_store[store_id] != row["service_store_code"]:
            raise ValueError(f"组织名单中的门店映射到多个服务店编码: {store_id}")
        code_by_store[store_id] = row["service_store_code"]
    result: list[str] = []
    seen_codes: set[str] = set()
    for poi_id in poi_ids:
        store_id = resolved[poi_id]
        code = code_by_store.get(store_id)
        if code is None:
            raise ValueError(f"资格名单POI未出现在组织名单中: {poi_id}")
        if code in seen_codes:
            raise ValueError(f"资格名单中的多个POI指向同一服务店编码: {code}")
        seen_codes.add(code)
        result.append(code)
    return result


def import_configuration_files(
    session: Session,
    organization_path: Path | None = None,
    eligibility_path: Path | None = None,
    now: datetime | None = None,
) -> dict[str, object]:
    """Validate roster files and publish one complete ranking configuration.

    Args:
        session: Caller-owned SQLAlchemy session. The caller commits the transaction.
        organization_path: Formal store organization XLSX or UTF-8 CSV file.
        eligibility_path: Applicable-store XLSX or UTF-8 CSV file.
        now: Server time override for deterministic tests. Defaults to current UTC time.

    Returns:
        Published versions, counts, actual effective times, and change status.

    Raises:
        ValueError: A file, POI mapping, version precondition, or row is invalid.
    """
    if organization_path is None and eligibility_path is None:
        raise ValueError("至少提供一份组织名单或资格名单")

    if session.get_bind().dialect.name == "postgresql":
        session.execute(
            text("SELECT pg_advisory_xact_lock(:key)"),
            {"key": RANKING_WRITE_LOCK},
        )

    latest_orgs, previous_org_version, previous_org_effective = _latest_organizations(session)
    latest_codes, previous_eligibility_version, previous_eligibility_effective = _latest_eligibility(
        session
    )
    has_complete_baseline = previous_org_version is not None and previous_eligibility_version is not None
    if not has_complete_baseline and (organization_path is None or eligibility_path is None):
        raise ValueError("首次导入必须同时提供组织名单和资格名单")

    organizations = (
        _parse_organizations(session, Path(organization_path))
        if organization_path is not None
        else latest_orgs
    )
    if eligibility_path is not None:
        eligibility_pois = _parse_eligibility_pois(Path(eligibility_path))
        eligible_codes = _eligible_codes(session, eligibility_pois, organizations)
    else:
        eligible_codes = latest_codes

    effective_from = DATA_START if not has_complete_baseline else _utc(now or datetime.now(timezone.utc))
    published = publish_configuration(
        session,
        organizations=organizations,
        eligible_codes=eligible_codes,
        effective_from=effective_from,
    )
    organization_changed = published["mapping_version"] != previous_org_version
    eligibility_changed = published["eligibility_version"] != previous_eligibility_version
    organization_effective = (
        effective_from if organization_changed else previous_org_effective
    )
    eligibility_effective = (
        effective_from if eligibility_changed else previous_eligibility_effective
    )
    if organization_effective is None or eligibility_effective is None:
        raise RuntimeError("publisher did not produce a complete ranking configuration")
    latest_effective = max(organization_effective, eligibility_effective)
    if not has_complete_baseline:
        change_status = "created"
    elif organization_changed or eligibility_changed:
        change_status = "updated"
    else:
        change_status = "unchanged"
    changed = organization_changed or eligibility_changed
    missing_organization_count = sum(
        any(
            row[field] == MISSING_ORGANIZATION_NAME
            for field in ("group_name", "district_name", "area_name")
        )
        for row in organizations
    )
    return {
        **published,
        "store_count": len(organizations),
        "eligible_store_count": len(eligible_codes),
        "eligibility_count": len(eligible_codes),
        "missing_organization_count": missing_organization_count,
        "effective_from": _utc(latest_effective).isoformat(),
        "organization_effective_from": _utc(organization_effective).isoformat(),
        "eligibility_effective_from": _utc(eligibility_effective).isoformat(),
        "organization_changed": organization_changed,
        "eligibility_changed": eligibility_changed,
        "changed": changed,
        "change_status": change_status,
    }
