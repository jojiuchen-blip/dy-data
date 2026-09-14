"""Excel ranking sheets: one metric per sheet, one section per selected level."""
from io import BytesIO
import json
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.api.dy_api.ranking_snapshots import read_snapshot_report, sort_ranking_rows

LEVEL_LABELS = {"group": "集团", "service_center": "中心", "district": "大区", "area": "区域", "store": "门店"}
METRIC_LABELS = {"order_average": "抖音店均订单量排行", "follow_24h_rate": "24小时有效跟进率排行", "verification_rate": "订单核销率排行"}
METRIC_COLUMNS = {
    "order_average": [("order_average", "店均订单量", "0.00"), ("order_count", "抖音订单量（辅助）", "0"), ("store_count", "适用门店数", "0")],
    "follow_24h_rate": [("follow_24h_rate", "24小时有效跟进率", "0.00%"), ("follow_rate", "不限24小时跟进率（辅助）", "0.00%"),
        ("follow_numerator", "24小时有效跟进轮次", "0"), ("follow_denominator", "正式分配轮次", "0"), ("follow_any_numerator", "不限24小时跟进轮次", "0")],
    "verification_rate": [("verification_rate", "订单核销率", "0.00%"), ("verification_numerator", "本店成功核销订单数", "0"), ("verification_denominator", "正式分配关联订单数", "0")],
}


def parse_export_choices(value: str, labels: dict[str, str]) -> list[str]:
    choices = [part.strip() for part in value.split(",")]
    if not choices or any(part not in labels for part in choices):
        raise ValueError("请至少选择一项有效的导出层级和指标")
    return [key for key in labels if key in choices]


def build_ranking_workbook(session, *, levels: list[str], metrics: list[str], **report_args) -> bytes:
    reports = {}
    for level in levels:
        report = read_snapshot_report(session, level=level, page_size=None, **report_args)
        report_args["run_id"] = report["snapshot_id"]  # Pin even preview reads to a single batch.
        reports[level] = report
    book = Workbook()
    book.remove(book.active)
    for metric in metrics:
        sheet = book.create_sheet(METRIC_LABELS[metric])
        columns = [("rank", "排名", "0"), ("name", "组织名称", "@"), ("key", "组织归属 / 门店编号", "@"), *METRIC_COLUMNS[metric]]
        width = len(columns)
        row_number = 0

        def append(values: list) -> int:
            nonlocal row_number
            row_number += 1
            sheet.append(values)
            return row_number

        def banner(text: str, *, strong: bool = False):
            row = append([text])
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
            cell = sheet.cell(row, 1)
            cell.font = Font(name="微软雅黑", bold=strong, size=12 if strong else 10, color="FFFFFF" if strong else "595959")
            cell.fill = PatternFill("solid", fgColor="C55A11" if strong else "FFF2E8")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[row].height = 30 if strong else 42

        first = reports[levels[0]]
        banner(METRIC_LABELS[metric], strong=True)
        banner(f"统计期间：{first['period_start']:%Y-%m-%d} 至 {first['period_end']:%Y-%m-%d}（北京时间）")
        observed = first["latest_observed_at"]
        if observed and observed.tzinfo is None:
            observed = observed.replace(tzinfo=ZoneInfo("UTC"))
        banner(f"数据更新时间：{observed.astimezone(ZoneInfo('Asia/Shanghai')):%Y-%m-%d %H:%M:%S}；快照：{first['snapshot_id']}" if observed else f"快照：{first['snapshot_id']}")
        filters = [(label, report_args.get(key)) for key, label in [("group_name", "集团"), ("service_center_name", "中心"), ("district_name", "大区"), ("area_name", "区域"), ("store_id", "门店")]]
        banner("数据范围：当前账号权限内；" + "；".join(f"{label}：{value}" for label, value in filters if value))
        banner(first["metric_definitions"][metric] + "。由高到低排列，同值并列，无样本不排名。")
        if first["preview_note"]:
            banner(first["preview_note"])
        if any(count > 0 and key != "follow_rounds_under_observation" for key, count in first["quality_json"].items()):
            banner("部分数据的门店归属或历史绑定资料尚待核验，当前排名仅供参考；未能唯一归属的数据暂未计入。")
        for level, report in reports.items():
            append([])
            banner(LEVEL_LABELS[level] + "排行", strong=True)
            header_row = append([label for _, label, _ in columns])
            for index in range(1, width + 1):
                cell = sheet.cell(header_row, index)
                cell.font = Font(name="微软雅黑", bold=True, color="9C3E00")
                cell.fill = PatternFill("solid", fgColor="FCE4D6")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[header_row].height = 32
            rows = sort_ranking_rows(report["rows"], metric)
            if not rows:
                banner("当前统计范围没有可展示数据")
            for item in rows:
                if level != "store":
                    try:
                        path = json.loads(item["key"])
                    except (ValueError, TypeError):
                        path = None
                    if isinstance(path, list) and all(isinstance(part, str) for part in path):
                        item["key"] = " / ".join(path)
                data_row = append([item.get(key) if item.get(key) is not None else ("—" if key == "rank" else "暂无样本") for key, _, _ in columns])
                for index, (_, _, fmt) in enumerate(columns, 1):
                    cell = sheet.cell(data_row, index)
                    # User-controlled organization names must always remain literal text.
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                    cell.number_format = fmt
                    cell.font = Font(name="微软雅黑", size=10)
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    if data_row % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="FFF8F3")
                sheet.row_dimensions[data_row].height = 30
        for index in range(1, width + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 10 if index == 1 else 28 if index in (2, 3) else 25
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
    output = BytesIO()
    book.save(output)
    return output.getvalue()
