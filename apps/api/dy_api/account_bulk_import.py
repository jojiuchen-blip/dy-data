"""Standard account onboarding workbook: parse, validate and export credentials."""
from io import BytesIO
from itertools import islice
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ['登录账号', '显示名称', '账号类型', '集团', '服务中心', '大区', '区域', '门店ID（多个用英文分号分隔）', '所属账户编号（选填）']
INITIAL_ACCOUNT_PASSWORD = '123456'
TYPES = {'最高管理员': None, '管理员': None, '集团账号': 'group', '服务中心账号': 'service_center', '大区账号': 'district', '区域账号': 'area', '门店账号': 'store'}


def _save(book):
    for sheet in book:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='9B461C')
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for column in 'ABCDEFGHI':
            sheet.column_dimensions[column].width = 25
        sheet.row_dimensions[1].height = 36
        if sheet.title == '填写说明':
            sheet.column_dimensions['B'].width = 95
            for number in range(2, sheet.max_row + 1):
                sheet.row_dimensions[number].height = 44
    output = BytesIO(); book.save(output)
    return output.getvalue()


def account_template(catalog=None):
    book = Workbook(); sheet = book.active; sheet.title = '账号开通'
    sheet.append(HEADERS)
    for row in range(2, 202):
        for column in (1, 8, 9):
            sheet.cell(row, column).number_format = '@'
    choices = DataValidation(type='list', formula1='"' + ','.join(TYPES) + '"', allow_blank=False)
    choices.errorTitle = '请选择账号类型'; choices.error = '请使用下拉列表中的账号类型'; choices.showErrorMessage = True
    sheet.add_data_validation(choices); choices.add('C2:C201')
    guide = book.create_sheet('填写说明')
    guide.append(['项目', '填写要求'])
    for values in [
        ('填写位置', '仅“账号开通”工作表会被导入，每行一个账号；其他工作表用于说明和查阅。每批最多200个账号。'),
        ('登录账号', '必填且唯一，建议手机号或英文账号；按文本填写，保留前导零。不要填写密码。'),
        ('显示名称', '必填，填写人员姓名或岗位名称。'),
        ('账号类型', '从下拉列表选择：最高管理员、管理员、集团账号、服务中心账号、大区账号、区域账号、门店账号。'),
        ('最高管理员/管理员', '拥有全部门店数据范围。组织管理人员请选对应组织账号类型，不要选全局管理员。'),
        ('集团账号', '填写集团；自动覆盖该集团下属门店。'),
        ('服务中心账号', '填写集团、服务中心。'),
        ('大区账号', '填写集团、服务中心、大区。'),
        ('区域账号', '填写集团、服务中心、大区、区域。必须填写完整路径以区分同名组织。'),
        ('门店账号', '填写系统门店ID，多个ID用英文分号分隔，例如00123;00456；组织列可留空。'),
        ('所属账户编号', '选填，如填写则必须唯一；不是门店ID。'),
        ('组织与门店名称', '必须与后台现有组织归属一致，可从“可选门店”工作表复制。无名单时从后台重新下载最新模板。'),
        ('开通流程', '最高管理员上传→校验预览→修正所有错误→确认批量创建。任一行无效则整批不创建。'),
        ('密码与结果', f'初始密码统一为{INITIAL_ACCOUNT_PASSWORD}，无需在表格填写。登录后请在账号菜单中选择“修改密码”。创建成功后可下载开通结果。'),
        ('重复提交', '已有账号会报错，不覆盖已有账号；如网络中断请先查账号列表，已创建账号通过重置密码处理。'),
        ('排行榜', '三个打榜指标对全部有效登录账号开放全量排名；线索和结算明细受账号组织范围限制。'),
    ]: guide.append(values)
    examples = book.create_sheet('示例（不导入）'); examples.append(HEADERS)
    for values in [
        ['group_zhang', '张经理', '集团账号', '示例集团', '', '', '', '', ''],
        ['center_li', '李经理', '服务中心账号', '示例集团', '示例中心', '', '', '', ''],
        ['district_wang', '王经理', '大区账号', '示例集团', '示例中心', '示例大区', '', '', ''],
        ['area_zhao', '赵经理', '区域账号', '示例集团', '示例中心', '示例大区', '示例区域', '', ''],
        ['store_qian', '钱店长', '门店账号', '', '', '', '', '00123;00456', ''],
    ]: examples.append(values)
    options = book.create_sheet('可选门店'); options.append(['门店ID', '门店名称', '集团', '服务中心', '大区', '区域'])
    for store in catalog or []:
        options.append([store.get(key, '') for key in ['store_id', 'store_name', 'group_name', 'service_center_name', 'district_name', 'area_name']])
    data = _save(book)
    return data


def read_account_rows(content: bytes, filename: str):
    if len(content) > 5 * 1024 * 1024 or not filename.lower().endswith('.xlsx'):
        raise ValueError('请上传不超过5MB的xlsx账号开通模板')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(row.file_size for row in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError('Excel内容过大，请使用标准模板')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            sheet = book['账号开通']; sheet.reset_dimensions()
            rows = list(islice(sheet.iter_rows(min_col=1, max_col=9, values_only=True), 202))
        finally: book.close()
    except (BadZipFile, KeyError, OSError, ParseError, TypeError) as exc:
        raise ValueError('无法读取“账号开通”工作表，请使用标准xlsx模板') from exc
    if not rows or list(rows[0]) != HEADERS:
        raise ValueError('表头不匹配，请使用最新账号开通模板，不要修改表头')
    if len(rows) > 201:
        raise ValueError('每批最多200个账号，请拆分表格')
    return [(number, row) for number, row in enumerate(rows[1:], 2) if any(value is not None and str(value).strip() for value in row)]


def credential_workbook(rows):
    book = Workbook(); sheet = book.active; sheet.title = '账号开通结果'
    sheet.append(['登录账号', '显示名称', '账号类型', '初始密码', '数据可见范围'])
    for row in rows: sheet.append(row)
    return _save(book)
