"""Readable account scope templates and non-mutating import previews."""
import csv
from io import BytesIO, StringIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from itertools import islice

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def store_template(catalog: list[dict]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = '门店名单'
    sheet.append(['门店ID', '门店名称'])
    sheet.freeze_panes = 'A2'
    sheet.column_dimensions['A'].width = 30
    sheet.column_dimensions['B'].width = 42
    for row in range(2, 1002):
        sheet.cell(row, 1).number_format = '@'
    instructions = book.create_sheet('填写说明')
    for line in [
        '1. 在“门店名单”填写，每行一家门店，第一行表头请保留。',
        '2. 门店ID必填，必须使用系统门店ID；不是POI ID、账号ID或服务门店编码。',
        '3. 门店ID为文本，保留前导零；请从“可选门店”复制ID，勿转成数字或科学计数法。',
        '4. 门店名称可选，用于核对；填写后必须与系统名称一致。',
        '5. 支持xlsx、UTF-8/GB18030 CSV或TXT；CSV第一列为门店ID，第二列为门店名称。',
        '6. 重复ID自动合并；导入先预览，错误行须修正后重新上传；应用后还需保存账号。',
    ]:
        instructions.append([line])
    instructions.column_dimensions['A'].width = 110
    options = book.create_sheet('可选门店')
    options.append(['门店ID', '门店名称', '组织归属'])
    for item in catalog:
        options.append([item['store_id'], item['store_name'], ' / '.join(item.get(k, '') for k in ('group_name', 'service_center_name', 'district_name', 'area_name'))])
        for cell in options[options.max_row]:
            cell.data_type = 's'
            cell.number_format = '@'
    for ws in (sheet, options):
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='FFF2E8')
        for col in ('A', 'B', 'C'):
            ws.column_dimensions[col].width = 38
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def preview_store_import(content: bytes, filename: str, catalog: list[dict]) -> dict:
    if len(content) > 5 * 1024 * 1024:
        raise ValueError('文件不能超过5MB')
    if filename.lower().endswith('.xlsx'):
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError('Excel解压后内容过大，请减少行数')
            book = load_workbook(BytesIO(content), read_only=True, data_only=False)
            sheet = book['门店名单'] if '门店名单' in book.sheetnames else book.active
            try:
                sheet.reset_dimensions()
                rows = list(islice(sheet.iter_rows(min_col=1, max_col=2, values_only=True), 10002))
            finally:
                book.close()
        except (BadZipFile, KeyError, OSError, ParseError, TypeError) as exc:
            raise ValueError("无法读取Excel，请使用下载的xlsx模板") from exc
    elif filename.lower().endswith(('.csv', '.txt')):
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                text = content.decode('gb18030')
            except UnicodeDecodeError as exc:
                raise ValueError('无法识别文本编码，请保存为UTF-8 CSV') from exc
        try:
            rows = list(islice(csv.reader(StringIO(text), delimiter='\t' if '\t' in text.split('\n')[0] else ','), 10002))
        except csv.Error as exc:
            raise ValueError('CSV格式不正确或单元格过长，请使用模板重新填写') from exc
    else:
        raise ValueError('请上传xlsx、csv或txt文件，旧版xls请另存为xlsx')
    if len(rows) > 10001:
        raise ValueError('每次最多导入10000行')
    by_id = {item['store_id']: item for item in catalog}
    ids, errors, seen, duplicate_count = [], [], set(), 0
    for number, row in enumerate(rows, 1):
        if not row or all(value is None or str(value).strip() == '' for value in row):
            continue
        raw = row[0]
        value = str(raw).strip() if raw is not None else ''
        if number == 1 and value.lower().replace(' ', '') in ('storeid', 'store_id', '门店id'):
            continue
        name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
        reason = ''
        if not isinstance(raw, str):
            reason = '门店ID必须为文本，请从可选门店复制，避免数字精度或前导零丢失'
        elif value not in by_id:
            reason = '门店ID不存在或无权分配，请核对可选门店名单'
        elif name and name != by_id[value]['store_name']:
            reason = '门店名称与ID不匹配'
        if reason:
            errors.append({'row': number, 'store_id': value, 'reason': reason})
        elif value in seen:
            duplicate_count += 1
        else:
            seen.add(value)
            ids.append(value)
    return {'store_ids': ids, 'stores': [by_id[value] for value in ids], 'errors': errors, 'duplicate_count': duplicate_count}
