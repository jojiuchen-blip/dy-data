from __future__ import annotations

from datetime import datetime, timedelta, timezone
from decimal import Decimal
from unittest.mock import Mock

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from apps.api.dy_api.models import (
    Base,
    ClueAllocationRule,
    ClueAllocationRuleVersion,
    ClueAssignmentRound,
    ClueCenterOrder,
    ClueFollowUpRecord,
    ClueMasterLead,
    ClueOrderStatusEvent,
    DataQualityIssue,
    DimStore,
    DimStorePoiMapping,
    RawDouyinClue,
    RawDouyinOrder,
    SettlementOrderDetail,
    StoreScoreSnapshot,
    StoreScoreSnapshotRun,
)
from apps.worker import clue_allocation
from apps.worker.repositories import upsert_data_quality_issue


def _dt(day: int, hour: int = 9) -> datetime:
    return datetime(2026, 7, day, hour, tzinfo=timezone.utc)


def _raw_clue(
    key: str,
    *,
    clue_id: str | None,
    order_id: str,
    order_status: str = "履约中",
    telephone: str | None = None,
    follow_poi_id: str | None = "poi-anchor",
    intention_poi_id: str | None = "poi-intention",
    auto_province_name: str | None = None,
    auto_city_name: str | None = None,
) -> RawDouyinClue:
    return RawDouyinClue(
        clue_row_key=key,
        clue_id=clue_id,
        order_id=order_id,
        order_status=order_status,
        telephone=telephone,
        follow_poi_id=follow_poi_id,
        intention_poi_id=intention_poi_id,
        auto_province_name=auto_province_name,
        auto_city_name=auto_city_name,
        create_time_detail=_dt(1),
        fetched_at=_dt(1),
        raw_payload={"clue_id": clue_id, "follow_poi_id": follow_poi_id},
        imported_at=_dt(1),
        updated_at=_dt(1),
    )


def _master_by_clue_id(session: Session, clue_id: str) -> ClueMasterLead | None:
    return session.scalar(select(ClueMasterLead).where(ClueMasterLead.canonical_clue_id == clue_id))


def _store(store_id: str, *, city_code: str, active: bool = True) -> DimStore:
    return DimStore(
        store_id=store_id,
        store_name=store_id,
        is_active=active,
        standard_province=city_code,
        standard_city=city_code,
        city_code=city_code,
        longitude=Decimal("121.470000"),
        latitude=Decimal("31.230000"),
        is_douyin_clue_applicable=True,
        participates_in_clue_allocation=True,
        location_source="test",
        location_status="valid",
        location_updated_at=_dt(1),
    )


def _published_score_rule_version(
    session: Session,
    *,
    rule_id: str,
    min_samples: int,
) -> ClueAllocationRuleVersion:
    rule = ClueAllocationRule(
        rule_id=rule_id,
        rule_name=rule_id,
        scope_type="global",
        scope_key=rule_id,
        created_by="test-admin",
    )
    version = ClueAllocationRuleVersion(
        rule_version_id=f"{rule_id}-v1",
        rule_id=rule.rule_id,
        version_no=1,
        status="published",
        auto_expiry_enabled=True,
        first_follow_up_sla_hours=24,
        protection_days=7,
        conversion_weight=Decimal("0.7"),
        follow_24h_weight=Decimal("0.3"),
        lookback_days=30,
        min_samples=min_samples,
        created_by="test-admin",
    )
    session.add_all([rule, version])
    session.commit()
    return version


def test_m1_clue_master_and_anchor_schema_is_declared() -> None:
    tables = Base.metadata.tables

    assert {"clue_master_leads", "clue_order_status_events"}.issubset(tables)
    assert {"follow_poi_id", "intention_poi_id"}.issubset(tables["raw_douyin_clues"].columns.keys())
    assert {
        "standard_province",
        "standard_city",
        "city_code",
        "longitude",
        "latitude",
        "is_douyin_clue_applicable",
        "participates_in_clue_allocation",
        "location_source",
        "location_status",
    }.issubset(tables["dim_stores"].columns.keys())
    assert {
        "lead_key",
        "source_clue_row_key",
        "source_identity_key",
        "order_id",
        "raw_order_status",
        "lifecycle_status",
        "pool_location",
        "allocation_state",
        "current_assignment_round_id",
        "allocation_cycle_id",
        "ended_without_assignment",
        "anchor_poi_id",
        "anchor_store_id",
        "anchor_source",
        "anchor_unavailable_reason",
    }.issubset(tables["clue_master_leads"].columns.keys())


def test_m1_score_schema_is_declared_without_reusing_legacy_rounds_as_formal_data() -> None:
    tables = Base.metadata.tables

    assert {"store_score_snapshot_runs", "store_score_snapshots"}.issubset(tables)
    assert {"execution_mode", "matured_at", "terminal_reason"}.issubset(
        tables["clue_assignment_rounds"].columns.keys()
    )
    assert {
        "snapshot_run_id",
        "snapshot_date",
        "store_id",
        "city_code",
        "conversion_numerator",
        "conversion_denominator",
        "follow_24h_numerator",
        "follow_24h_denominator",
        "conversion_value_source",
        "follow_24h_value_source",
        "composite_score",
    }.issubset(tables["store_score_snapshots"].columns.keys())


def test_data_quality_issue_upsert_can_defer_flush() -> None:
    session = Mock(spec=Session)
    session.new = ()

    upsert_data_quality_issue(
        session,
        "deferred-quality-issue",
        issue_type="clue_anchor_unavailable",
        message="anchor unavailable",
        raw_context_json={"reason": "follow_poi_missing"},
        flush=False,
    )

    session.flush.assert_not_called()


def test_data_quality_issue_upsert_deduplicates_pending_issue_before_flush(db_session: Session) -> None:
    for _ in range(2):
        upsert_data_quality_issue(
            db_session,
            "deferred-duplicate-quality-issue",
            issue_type="clue_anchor_unavailable",
            message="anchor unavailable",
            raw_context_json={"reason": "follow_poi_missing"},
            flush=False,
        )

    db_session.flush()

    assert db_session.scalar(select(func.count()).select_from(DataQualityIssue)) == 1


def test_materialize_master_leads_keeps_terminal_clues_without_assignment_rounds(
    db_session: Session,
) -> None:
    db_session.add_all(
        [
            _store("store-anchor", city_code="上海"),
            DimStorePoiMapping(store_id="store-anchor", poi_id="poi-anchor", mapping_source="test"),
            _raw_clue("active-row", clue_id="active-clue", order_id="active-order"),
            _raw_clue("verified-row", clue_id="verified-clue", order_id="verified-order", order_status="已核销"),
            _raw_clue("refund-row", clue_id="refund-clue", order_id="refund-order", order_status="已退款"),
        ]
    )
    db_session.commit()

    stats = clue_allocation.materialize_clue_master_leads(db_session, now=_dt(2))

    assert stats == {"master_leads": 3, "closed_leads": 2, "headquarters_pool": 0}
    active = _master_by_clue_id(db_session, "active-clue")
    verified = _master_by_clue_id(db_session, "verified-clue")
    refunded = _master_by_clue_id(db_session, "refund-clue")
    assert active is not None
    assert active.lifecycle_status == "active"
    assert active.pool_location is None
    assert active.allocation_state == "pending_allocation"
    assert active.ended_without_assignment is False
    assert active.anchor_store_id == "store-anchor"
    assert verified is not None
    assert verified.lifecycle_status == "closed_verified"
    assert verified.pool_location == "closed"
    assert verified.allocation_state == "closed"
    assert verified.ended_without_assignment is True
    assert refunded is not None
    assert refunded.lifecycle_status == "closed_refunded"
    assert refunded.pool_location == "closed"
    assert refunded.allocation_state == "closed"
    assert refunded.ended_without_assignment is True
    assert db_session.scalar(select(func.count()).select_from(ClueAssignmentRound)) == 0
    assert db_session.scalar(select(func.count()).select_from(ClueOrderStatusEvent)) == 3

    clue_allocation.materialize_clue_master_leads(db_session, now=_dt(3))
    assert db_session.scalar(select(func.count()).select_from(ClueOrderStatusEvent)) == 3


def test_master_lead_uses_follow_poi_only_and_routes_invalid_anchor_to_headquarters(
    db_session: Session,
) -> None:
    db_session.add(_raw_clue("missing-anchor", clue_id="missing-anchor", order_id="missing-order", follow_poi_id=None))
    db_session.commit()

    stats = clue_allocation.materialize_clue_master_leads(db_session, now=_dt(2))

    assert stats == {"master_leads": 1, "closed_leads": 0, "headquarters_pool": 1}
    master = _master_by_clue_id(db_session, "missing-anchor")
    assert master is not None
    assert master.anchor_poi_id is None
    assert master.anchor_store_id is None
    assert master.pool_location == "headquarters_pool"
    assert master.allocation_state == "headquarters"
    assert master.anchor_unavailable_reason == "follow_poi_missing"


def test_materialization_maps_order_api_waiting_use_status_to_active(db_session: Session) -> None:
    db_session.add_all(
        [
            _store("store-anchor", city_code="上海"),
            DimStorePoiMapping(store_id="store-anchor", poi_id="poi-anchor", mapping_source="test"),
            _raw_clue(
                "waiting-use-row",
                clue_id="waiting-use-clue",
                order_id="waiting-use-order",
                order_status="交易关闭",
            ),
            RawDouyinOrder(
                order_id="waiting-use-order",
                order_status="201",
                raw_payload={"order_status": "201"},
                created_at=_dt(1),
                updated_at=_dt(1),
            ),
        ]
    )
    db_session.commit()

    clue_allocation.materialize_clue_master_leads(db_session, now=_dt(2))

    master = _master_by_clue_id(db_session, "waiting-use-clue")
    assert master is not None
    assert master.normalized_order_status == "active"
    assert master.lifecycle_status == "active"
    assert master.allocation_state == "pending_allocation"


def test_materialization_persists_new_master_before_creating_headquarters_entry(
    db_session: Session,
) -> None:
    db_session.execute(text("PRAGMA foreign_keys = ON"))
    db_session.commit()
    db_session.add(_raw_clue("fk-anchor", clue_id="fk-anchor", order_id="fk-order", follow_poi_id=None))
    db_session.commit()

    stats = clue_allocation.materialize_clue_master_leads(db_session, now=_dt(2))

    assert stats == {"master_leads": 1, "closed_leads": 0, "headquarters_pool": 1}


def test_master_lead_merges_missing_clue_id_when_contact_and_order_are_later_available(
    db_session: Session,
) -> None:
    db_session.add_all(
        [
            _raw_clue(
                "raw-before-clue-id",
                clue_id=None,
                order_id="identity-order",
                telephone="13812345678",
            ),
            _raw_clue(
                "row-after-clue-id",
                clue_id="later-clue-id",
                order_id="identity-order",
                telephone="13812345678",
            ),
        ]
    )
    db_session.commit()

    stats = clue_allocation.materialize_clue_master_leads(db_session, now=_dt(2))

    assert stats["master_leads"] == 1
    master = _master_by_clue_id(db_session, "later-clue-id")
    assert master is not None
    assert master.source_clue_row_key == "raw-before-clue-id"
    assert db_session.scalar(select(func.count()).select_from(ClueMasterLead)) == 1

    db_session.add(
        _raw_clue(
            "raw-after-clue-id",
            clue_id=None,
            order_id="identity-order",
            telephone="13812345678",
        )
    )
    db_session.commit()
    clue_allocation.materialize_clue_master_leads(db_session, now=_dt(3))

    assert _master_by_clue_id(db_session, "later-clue-id") is not None


def test_terminal_master_status_closes_legacy_current_round(db_session: Session) -> None:
    db_session.add_all(
        [
            _raw_clue("terminal-row", clue_id="terminal-clue", order_id="terminal-order", order_status="已核销"),
            ClueCenterOrder(
                order_id="terminal-order",
                lead_status="active",
                current_assignment_round_id="terminal-order-1",
                current_round_no=1,
                current_round_status="active_unfollowed",
                created_at=_dt(1),
                updated_at=_dt(1),
            ),
            ClueAssignmentRound(
                assignment_round_id="terminal-order-1",
                order_id="terminal-order",
                round_no=1,
                round_status="active_unfollowed",
                execution_mode="legacy",
                created_at=_dt(1),
                updated_at=_dt(1),
            ),
        ]
    )
    db_session.commit()

    clue_allocation.materialize_clue_master_leads(db_session, now=_dt(2))

    round_row = db_session.get(ClueAssignmentRound, "terminal-order-1")
    order = db_session.get(ClueCenterOrder, "terminal-order")
    assert round_row is not None
    assert round_row.round_status == "closed_order_verified"
    assert round_row.terminal_reason == "order_verified"
    assert order is not None
    assert order.lead_status == "converted"
    assert order.current_round_status == "closed_order_verified"


def test_imported_store_locations_use_poi_mapping_and_candidate_eligibility(
    db_session: Session, tmp_path
) -> None:
    db_session.add_all(
        [
            DimStore(store_id="store-open", store_name="Open", is_active=True),
            DimStore(store_id="store-closed", store_name="Closed", is_active=True),
            DimStorePoiMapping(store_id="store-open", poi_id="6601132611886647299", mapping_source="test"),
            DimStorePoiMapping(store_id="store-closed", poi_id="6601133876418971651", mapping_source="test"),
        ]
    )
    db_session.commit()
    workbook_path = tmp_path / "store-locations.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["门店ID", "门店名称", "服务店名称", "状态备注", "服务店代码", "经度", "纬度", "门店所在省份", "门店所在城市"])
    sheet.append(["6601132611886647299", "Open", "Open Service", "", "A", 121.47, 31.23, "上海", "上海"])
    sheet.append(["6601133876418971651", "Closed", "Closed Service", "关闭", "B", 121.48, 31.24, "上海", "上海"])
    workbook.save(workbook_path)

    stats = clue_allocation.import_store_locations(
        db_session,
        workbook_path,
        enable_participation=True,
        now=_dt(2),
    )

    assert stats == {"rows": 2, "updated": 2, "unmapped": 0, "invalid": 0}
    open_store = db_session.get(DimStore, "store-open")
    closed_store = db_session.get(DimStore, "store-closed")
    assert open_store is not None
    assert open_store.city_code == "上海"
    assert open_store.standard_province == "上海"
    assert open_store.is_douyin_clue_applicable is True
    assert open_store.participates_in_clue_allocation is True
    assert closed_store is not None
    assert closed_store.is_douyin_clue_applicable is False
    assert closed_store.location_status == "closed"
    assert [store.store_id for store in clue_allocation.eligible_candidate_stores(db_session, city_code="上海")] == [
        "store-open"
    ]
    assert clue_allocation.haversine_km(31.23, 121.47, 31.23, 121.47) == pytest.approx(0)


def test_raw_anchor_evidence_completes_missing_store_province_before_candidate_evaluation(
    db_session: Session,
) -> None:
    db_session.add_all(
        [
            DimStore(
                store_id="store-needs-province",
                store_name="Needs Province",
                is_active=True,
                standard_city="上海市",
                city_code="上海",
                longitude=Decimal("121.470000"),
                latitude=Decimal("31.230000"),
                is_douyin_clue_applicable=False,
                participates_in_clue_allocation=True,
                location_status="partial",
            ),
            DimStorePoiMapping(store_id="store-needs-province", poi_id="poi-needs-province", mapping_source="test"),
            _raw_clue(
                "province-evidence-row",
                clue_id="province-evidence-clue",
                order_id="province-evidence-order",
                follow_poi_id="poi-needs-province",
                auto_province_name="上海市",
                auto_city_name="上海市",
            ),
        ]
    )
    db_session.commit()

    clue_allocation.materialize_clue_master_leads(db_session, now=_dt(2))

    store = db_session.get(DimStore, "store-needs-province")
    assert store is not None
    assert store.standard_province == "上海市"
    assert store.location_status == "valid"
    assert store.is_douyin_clue_applicable is True
    assert [candidate.store_id for candidate in clue_allocation.eligible_candidate_stores(db_session)] == [
        "store-needs-province"
    ]


def test_location_import_uses_raw_poi_evidence_before_enabling_candidate_participation(
    db_session: Session,
    tmp_path,
) -> None:
    db_session.add_all(
        [
            DimStore(store_id="store-import-evidence", store_name="Import Evidence", is_active=True),
            DimStorePoiMapping(store_id="store-import-evidence", poi_id="poi-import-evidence", mapping_source="test"),
            _raw_clue(
                "import-evidence-row",
                clue_id="import-evidence-clue",
                order_id="import-evidence-order",
                follow_poi_id="poi-import-evidence",
                auto_province_name="上海市",
                auto_city_name="上海市",
            ),
        ]
    )
    db_session.commit()
    workbook_path = tmp_path / "store-locations-without-province.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["门店ID", "经度", "纬度", "门店所在城市"])
    sheet.append(["poi-import-evidence", 121.47, 31.23, "上海市"])
    workbook.save(workbook_path)

    clue_allocation.import_store_locations(
        db_session,
        workbook_path,
        enable_participation=True,
        now=_dt(2),
    )

    store = db_session.get(DimStore, "store-import-evidence")
    assert store is not None
    assert store.standard_province == "上海市"
    assert store.location_status == "valid"
    assert store.participates_in_clue_allocation is True


def test_score_snapshots_use_formal_mature_rounds_and_city_global_fallbacks(
    db_session: Session,
) -> None:
    store_a = _store("store-a", city_code="上海")
    store_b = _store("store-b", city_code="上海")
    store_c = _store("store-c", city_code="北京")
    assigned_at = _dt(1)
    db_session.add_all(
        [
            store_a,
            store_b,
            store_c,
            ClueAssignmentRound(
                assignment_round_id="formal-1",
                order_id="formal-order-1",
                round_no=1,
                assigned_store_id="store-a",
                assigned_at=assigned_at,
                round_status="closed_order_verified",
                execution_mode="formal",
                matured_at=_dt(3),
                created_at=assigned_at,
                updated_at=_dt(2),
            ),
            ClueAssignmentRound(
                assignment_round_id="formal-2",
                order_id="formal-order-2",
                round_no=1,
                assigned_store_id="store-a",
                assigned_at=assigned_at,
                round_status="expired",
                execution_mode="formal",
                matured_at=_dt(2),
                terminal_reason="sla_expired",
                created_at=assigned_at,
                updated_at=_dt(2),
            ),
            ClueAssignmentRound(
                assignment_round_id="formal-3-completed-early",
                order_id="formal-order-3",
                round_no=1,
                assigned_store_id="store-a",
                assigned_at=assigned_at,
                round_status="closed_order_verified",
                execution_mode="formal",
                matured_at=assigned_at + timedelta(hours=1),
                terminal_reason="order_verified",
                created_at=assigned_at,
                updated_at=assigned_at + timedelta(hours=1),
            ),
            ClueAssignmentRound(
                assignment_round_id="legacy-ignored",
                order_id="legacy-order",
                round_no=1,
                assigned_store_id="store-c",
                assigned_at=assigned_at,
                round_status="closed_order_verified",
                execution_mode="legacy",
                matured_at=_dt(2),
                created_at=assigned_at,
                updated_at=_dt(2),
            ),
            SettlementOrderDetail(
                coupon_id="coupon-formal-1",
                order_id="formal-order-1",
                product_type="test",
                sale_time=assigned_at,
                is_verified=True,
                verify_store_id="store-a",
                verify_store_name="store-a",
                verify_time=_dt(3),
                relation_type="same_store",
                is_commissionable=False,
                is_refund_excluded=False,
                paid_amount_cent=100,
                commission_rate=Decimal("0"),
                receivable_commission_cent=0,
                payable_commission_cent=0,
                updated_at=_dt(2),
            ),
            SettlementOrderDetail(
                coupon_id="coupon-formal-3",
                order_id="formal-order-3",
                product_type="test",
                sale_time=assigned_at,
                is_verified=True,
                verify_store_id="store-a",
                verify_store_name="store-a",
                verify_time=assigned_at + timedelta(hours=1),
                relation_type="same_store",
                is_commissionable=False,
                is_refund_excluded=False,
                paid_amount_cent=100,
                commission_rate=Decimal("0"),
                receivable_commission_cent=0,
                payable_commission_cent=0,
                updated_at=assigned_at + timedelta(hours=1),
            ),
            ClueFollowUpRecord(
                follow_up_record_id="follow-formal-1",
                order_id="formal-order-1",
                assignment_round_id="formal-1",
                round_no=1,
                assigned_store_id="store-a",
                follow_result="appointment",
                created_at=assigned_at + timedelta(hours=12),
            ),
            ClueFollowUpRecord(
                follow_up_record_id="follow-late",
                order_id="formal-order-2",
                assignment_round_id="formal-2",
                round_no=1,
                assigned_store_id="store-a",
                follow_result="unreachable",
                created_at=assigned_at + timedelta(hours=25),
            ),
        ]
    )
    db_session.commit()
    rule_version = _published_score_rule_version(db_session, rule_id="metrics-rule", min_samples=2)

    result = clue_allocation.refresh_store_score_snapshots(
        db_session,
        rule_version_id=rule_version.rule_version_id,
        now=_dt(10),
        run_mode="manual",
    )

    assert result["snapshots"] == 3
    rows = {
        row.store_id: row
        for row in db_session.scalars(select(StoreScoreSnapshot).where(StoreScoreSnapshot.snapshot_run_id == result["snapshot_run_id"])).all()
    }
    run = db_session.get(StoreScoreSnapshotRun, result["snapshot_run_id"])
    assert run is not None
    assert run.snapshot_count == 3
    assert rows["store-a"].conversion_rate == Decimal("0.666667")
    assert rows["store-a"].follow_24h_rate == Decimal("0.500000")
    assert rows["store-a"].follow_24h_denominator == 2
    assert rows["store-a"].conversion_value_source == "store"
    assert rows["store-b"].conversion_value_source == "city"
    assert rows["store-b"].follow_24h_value_source == "city"
    assert rows["store-c"].conversion_value_source == "global"
    assert rows["store-c"].follow_24h_value_source == "global"


def test_score_attributes_verification_to_the_latest_formal_round_only(db_session: Session) -> None:
    assigned_first = _dt(1)
    assigned_second = _dt(2)
    db_session.add_all(
        [
            _store("store-first", city_code="上海"),
            _store("store-second", city_code="上海"),
            ClueAssignmentRound(
                assignment_round_id="same-order-first",
                order_id="same-order",
                round_no=1,
                assigned_store_id="store-first",
                assigned_at=assigned_first,
                round_status="expired",
                execution_mode="formal",
                matured_at=assigned_second,
                terminal_reason="reassigned",
                created_at=assigned_first,
                updated_at=assigned_second,
            ),
            ClueAssignmentRound(
                assignment_round_id="same-order-second",
                order_id="same-order",
                round_no=2,
                assigned_store_id="store-second",
                assigned_at=assigned_second,
                round_status="closed_order_verified",
                execution_mode="formal",
                matured_at=_dt(3),
                terminal_reason="order_verified",
                created_at=assigned_second,
                updated_at=_dt(3),
            ),
            SettlementOrderDetail(
                coupon_id="coupon-same-order",
                order_id="same-order",
                product_type="test",
                sale_time=assigned_first,
                is_verified=True,
                verify_store_id="store-second",
                verify_store_name="store-second",
                verify_time=_dt(3),
                relation_type="same_store",
                is_commissionable=False,
                is_refund_excluded=False,
                paid_amount_cent=100,
                commission_rate=Decimal("0"),
                receivable_commission_cent=0,
                payable_commission_cent=0,
                updated_at=_dt(3),
            ),
        ]
    )
    db_session.commit()
    rule_version = _published_score_rule_version(db_session, rule_id="verification-rule", min_samples=1)

    result = clue_allocation.refresh_store_score_snapshots(
        db_session,
        rule_version_id=rule_version.rule_version_id,
        now=_dt(5),
        run_mode="manual",
    )

    rows = {
        row.store_id: row
        for row in db_session.scalars(
            select(StoreScoreSnapshot).where(StoreScoreSnapshot.snapshot_run_id == result["snapshot_run_id"])
        ).all()
    }
    assert rows["store-first"].conversion_numerator == 0
    assert rows["store-first"].conversion_denominator == 1
    assert rows["store-second"].conversion_numerator == 1
    assert rows["store-second"].conversion_denominator == 1


def test_scheduled_score_refresh_skips_without_a_rule_version_after_shanghai_three_am(db_session: Session) -> None:
    db_session.add(_store("scheduled-store", city_code="上海"))
    db_session.commit()

    before = clue_allocation.refresh_due_store_score_snapshots(
        db_session,
        now=datetime(2026, 7, 10, 18, 30, tzinfo=timezone.utc),
    )
    after = clue_allocation.refresh_due_store_score_snapshots(
        db_session,
        now=datetime(2026, 7, 10, 19, 1, tzinfo=timezone.utc),
    )
    repeat = clue_allocation.refresh_due_store_score_snapshots(
        db_session,
        now=datetime(2026, 7, 10, 20, 0, tzinfo=timezone.utc),
    )

    assert before == {"snapshot_run_id": None, "snapshots": 0, "skipped": "before_schedule"}
    assert after == {"snapshot_run_id": None, "snapshots": 0, "skipped": "no_rule_versions"}
    assert repeat == {"snapshot_run_id": None, "snapshots": 0, "skipped": "no_rule_versions"}


def test_scheduled_score_refresh_does_not_record_an_unbound_empty_run(db_session: Session) -> None:
    first = clue_allocation.refresh_due_store_score_snapshots(
        db_session,
        now=datetime(2026, 7, 10, 19, 1, tzinfo=timezone.utc),
    )
    repeat = clue_allocation.refresh_due_store_score_snapshots(
        db_session,
        now=datetime(2026, 7, 10, 20, 0, tzinfo=timezone.utc),
    )

    assert first == {"snapshot_run_id": None, "snapshots": 0, "skipped": "no_rule_versions"}
    assert repeat == {"snapshot_run_id": None, "snapshots": 0, "skipped": "no_rule_versions"}
    assert db_session.scalar(select(func.count()).select_from(StoreScoreSnapshotRun)) == 0


def test_master_materialization_skips_when_its_transaction_lock_is_unavailable(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    db_session.add_all(
        [
            _store("locked-anchor", city_code="上海"),
            DimStorePoiMapping(store_id="locked-anchor", poi_id="locked-poi", mapping_source="test"),
            _raw_clue("locked-raw", clue_id="locked-clue", order_id="locked-order", follow_poi_id="locked-poi"),
        ]
    )
    db_session.commit()
    monkeypatch.setattr(clue_allocation, "_try_transaction_lock", lambda *args, **kwargs: False, raising=False)

    result = clue_allocation.materialize_clue_master_leads(db_session, now=_dt(2))

    assert result == {"master_leads": 0, "closed_leads": 0, "headquarters_pool": 0, "skipped": "locked"}
    assert db_session.scalar(select(func.count()).select_from(ClueMasterLead)) == 0


def test_scheduled_score_refresh_skips_when_its_transaction_lock_is_unavailable(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    db_session.add(_store("locked-score-store", city_code="上海"))
    db_session.commit()
    monkeypatch.setattr(clue_allocation, "_try_transaction_lock", lambda *args, **kwargs: False, raising=False)

    result = clue_allocation.refresh_due_store_score_snapshots(
        db_session,
        now=datetime(2026, 7, 10, 19, 1, tzinfo=timezone.utc),
    )

    assert result == {"snapshot_run_id": None, "snapshots": 0, "skipped": "locked"}
    assert db_session.scalar(select(func.count()).select_from(StoreScoreSnapshotRun)) == 0
