"""Exercise the real route/auth/serialization against synthetic snapshots."""
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, func

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "apps" / "api"))
from dy_api.main import create_app
from dy_api.auth import AuthContext, get_current_user
from dy_api.routes._data import get_session_dependency
from apps.api.dy_api.ranking_schema_v1 import metadata, runs
from apps.api.dy_api.ranking_snapshots import calculate_snapshot
from apps.worker.ranking_preview_fixture import seed_preview, START, END, CUTOFF, ELIGIBILITY_VERSION


@pytest.fixture()
def preview_client(db_session, monkeypatch):
    monkeypatch.setenv("DY_API_TEST_MODE", "true")
    monkeypatch.setenv("DY_SUPER_ADMIN_USERNAME", "ranking_test_admin")
    monkeypatch.setenv("DY_TEST_ADMIN_PASSWORD", "synthetic-test-only")
    monkeypatch.setenv("DY_SESSION_COOKIE_SECURE", "false")
    metadata.create_all(db_session.bind)
    seed_preview(db_session)
    calculate_snapshot(db_session, run_id="baseline", period_start=START, period_end=END,
                       observed_through=CUTOFF, roster_at=START, eligibility_version=ELIGIBILITY_VERSION)
    db_session.commit()
    app = create_app()
    app.state.ranking_snapshot_preview = True

    def session_override():
        yield db_session

    app.dependency_overrides[get_session_dependency] = session_override
    return TestClient(app)


PARAMS = {"periodStart": "2026-09-01", "periodEnd": "2026-09-30", "level": "store"}


def login(client):
    assert client.post("/api/v1/auth/login", json={"username": "ranking_test_admin",
        "password": "synthetic-test-only"}).status_code == 200


def test_preview_route_reads_snapshot_not_legacy_table(preview_client):
    login(preview_client)
    response = preview_client.get("/api/v1/dashboard/douyin-ranking", params=PARAMS)
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["dataMode"] == "synthetic"
    assert data["snapshotId"] == "baseline"
    assert data["totals"]["storeCount"] == 3
    assert data["totals"]["orderAverage"] == 3.333333
    assert data["totals"]["follow24hRate"] == .3
    assert data["totals"]["verificationRate"] == .5
    assert "虚拟测试" in data["previewNote"]


def test_missing_period_is_explicit_error_without_mock_fallback(preview_client):
    login(preview_client)
    response = preview_client.get("/api/v1/dashboard/douyin-ranking",
        params={**PARAMS, "periodStart": "2026-08-01", "periodEnd": "2026-08-31"})
    assert response.status_code == 422


def test_preview_still_requires_login(preview_client):
    assert preview_client.get("/api/v1/dashboard/douyin-ranking", params=PARAMS).status_code == 401


def test_preview_ranking_is_global_for_store_accounts(preview_client):
    auth = AuthContext(user_id=None, username="store-a", display_name="A", role="store",
        store_ids=("A",), auth_type="env_admin", store_scope_mode="explicit", page_keys=("A03",))
    preview_client.app.dependency_overrides[get_current_user] = lambda: auth
    response = preview_client.get("/api/v1/dashboard/douyin-ranking", params=PARAMS)
    assert response.status_code == 200
    data = response.json()["data"]
    assert {row["key"] for row in data["rows"]} == {"A", "B", "C"}
    assert data["totals"]["orderCount"] == 10
    assert data["totals"]["verificationRate"] == .5


def test_snapshot_read_never_creates_another_batch(preview_client, db_session):
    login(preview_client)
    before = db_session.scalar(select(func.count()).select_from(runs))
    for _ in range(2):
        assert preview_client.get("/api/v1/dashboard/douyin-ranking", params=PARAMS).status_code == 200
    assert db_session.scalar(select(func.count()).select_from(runs)) == before


EXPORT = "/api/v1/dashboard/douyin-ranking/export"


def workbook(response):
    from io import BytesIO
    from openpyxl import load_workbook
    assert response.status_code == 200, response.text[:200] if response.status_code != 200 else ""
    assert "spreadsheetml" in response.headers["content-type"]
    return load_workbook(BytesIO(response.content))


def test_export_three_metric_sheets_with_two_level_sections(preview_client):
    login(preview_client)
    book = workbook(preview_client.get(EXPORT, params={**PARAMS, "levels": "district,area"}))
    assert book.sheetnames == ["抖音店均订单量排行", "24小时有效跟进率排行", "订单核销率排行"]
    for sheet in book:
        values = [cell.value for row in sheet for cell in row]
        assert "大区排行" in values and "区域排行" in values
        assert "baseline" in str(values) and "2026-09-01" in str(values)
    assert "抖音订单量（辅助）" in [c.value for row in book.worksheets[0] for c in row]
    assert "不限24小时跟进率（辅助）" in [c.value for row in book.worksheets[1] for c in row]


def test_export_rankings_match_board_and_include_all_rows(preview_client):
    login(preview_client)
    book = workbook(preview_client.get(EXPORT, params={**PARAMS, "levels": "store", "pageSize": 1}))
    for sheet, metric in zip(book, ["order_average", "follow_24h_rate", "verification_rate"]):
        board = preview_client.get("/api/v1/dashboard/douyin-ranking", params={**PARAMS, "sortBy": metric}).json()["data"]
        export_rows = [row for row in sheet.iter_rows(values_only=True) if isinstance(row[0], int) or row[0] == "—"]
        assert len(export_rows) == board["total"]
        assert [(row[0], row[1]) for row in export_rows] == [(row["rank"] if row["rank"] is not None else "—", row["name"]) for row in board["rows"]]


@pytest.mark.parametrize("params", [{"levels": ""}, {"levels": "invalid"}, {"metrics": "order_count"},
    {"metrics": ""}, {"periodEnd": "2026-08-01"}, {"periodEnd": "2028-01-01"}])
def test_export_rejects_invalid_options(preview_client, params):
    login(preview_client)
    assert preview_client.get(EXPORT, params={**PARAMS, **params}).status_code == 422


def test_export_requires_login_and_allows_global_ranking(preview_client):
    assert preview_client.get(EXPORT, params=PARAMS).status_code == 401
    auth = AuthContext(user_id=None, username="store-a", display_name="A", role="store",
        store_ids=("A",), auth_type="env_admin", store_scope_mode="explicit", page_keys=("A03",))
    preview_client.app.dependency_overrides[get_current_user] = lambda: auth
    book = workbook(preview_client.get(EXPORT, params={**PARAMS, "levels": "store", "metrics": "order_average"}))
    rows = [row for row in book.active.iter_rows(values_only=True) if isinstance(row[0], int)]
    assert len(rows) == 3
    empty = workbook(preview_client.get(EXPORT, params={**PARAMS, "levels": "store", "storeId": "B"}))
    assert any(isinstance(row[0], int) and row[2] == "B" for row in empty.active.iter_rows(values_only=True))


def test_export_over_200_rows_and_formula_like_names_are_literal(preview_client, db_session):
    from apps.api.dy_api.ranking_schema_v1 import org_history, snapshots
    original = dict(db_session.execute(select(org_history).where(org_history.c.store_id == "A")).mappings().first())
    facts = [dict(row) for row in db_session.execute(select(snapshots).where(snapshots.c.store_id == "A")).mappings()]
    for index in range(205):
        store_id = f"extra-{index}"
        db_session.execute(org_history.insert().values(**{**original, "store_id": store_id,
            "service_store_code": store_id, "store_name": "=1+1" if index == 0 else store_id}))
        db_session.execute(snapshots.insert(), [{**row, "store_id": store_id} for row in facts])
    db_session.commit()
    login(preview_client)
    book = workbook(preview_client.get(EXPORT, params={**PARAMS, "levels": "store", "metrics": "order_average"}))
    rows = [row for row in book.active if isinstance(row[0].value, int)]
    assert len(rows) == 208
    formula_name = next(row[1] for row in rows if row[1].value == "=1+1")
    assert formula_name.data_type == "s"
    assert all(cell.data_type != "f" for row in book.active for cell in row)


def test_shared_metric_ranking_ties_nulls_and_auxiliary_values():
    from apps.api.dy_api.ranking_snapshots import sort_ranking_rows
    rows = [dict(name="A", order_average=2, order_count=100, follow_24h_rate=.1, follow_rate=.9, verification_rate=None),
            dict(name="B", order_average=4, order_count=8, follow_24h_rate=.5, follow_rate=.6, verification_rate=.2),
            dict(name="C", order_average=4, order_count=12, follow_24h_rate=None, follow_rate=None, verification_rate=.8)]
    for metric, names in [("order_average", ["B", "C", "A"]), ("follow_24h_rate", ["B", "A", "C"]), ("verification_rate", ["C", "B", "A"])]:
        ranked = sort_ranking_rows(rows, metric)
        assert [row["name"] for row in ranked] == names
    assert [row["rank"] for row in sort_ranking_rows(rows, "order_average")] == [1, 1, 3]
    assert sort_ranking_rows(rows, "follow_24h_rate")[-1]["rank"] is None
    assert "rank" not in rows[0]
