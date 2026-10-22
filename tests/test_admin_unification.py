"""DYDATA-105 real-route administrator and account scope regressions."""

from io import BytesIO

from openpyxl import load_workbook
import pytest
from sqlalchemy import func, select

from test_api_access_control import client, _login
from test_account_bulk_import import workbook
from apps.api.dy_api.models import (
    AccountPermissionAuditLog, DimStore, JobRun, User, UserPagePermissionOverride, UserStoreScope,
)
from dy_api.auth import hash_password_pbkdf2
from dy_api.access_control import ALL_PAGE_KEYS


def scoped_admin(db_session, role="admin"):
    user = db_session.get(User, "db-admin")
    user.role = role
    user.store_scope_mode = "specified"
    db_session.add(UserStoreScope(user_id=user.user_id, store_id="store-1"))
    db_session.commit()
    return user


def payload(username="created", **changes):
    return {
        "username": username, "display_name": username, "role": "store", "status": "active",
        "store_scope_mode": "specified", "store_ids": ["store-1"],
        "password": "secret", "password_confirm": "secret", **changes,
    }


@pytest.mark.parametrize("legacy_role", ["admin", "highest_admin", "viewer"])
def test_admin_fixed_pages_never_widen_persisted_scope(client, db_session, legacy_role):
    scoped_admin(db_session, legacy_role)
    db_session.add(UserPagePermissionOverride(user_id="db-admin", page_key="D02", effect="deny"))
    db_session.commit()
    _login(client, "db-admin")
    identity = client.get("/api/v1/auth/me").json()["data"]
    assert identity["role"] == "admin"
    assert identity["page_keys"] == list(ALL_PAGE_KEYS)
    assert identity["store_scope_mode"] == "specified"
    assert identity["store_ids"] == ["store-1"]
    assert identity["can_manage_accounts"] is True
    assert identity["can_manage_global_settings"] is False
    assert identity["is_highest_admin"] is False
    assert client.get("/api/v1/admin/access-control").status_code == 200
    assert client.get("/api/v1/admin/sync").status_code == 200


def test_scoped_account_management_requires_the_whole_target_scope(client, db_session):
    scoped_admin(db_session)
    for identifier, ids in (("outside", ["store-2"]), ("mixed", ["store-1", "store-2"]), ("empty", [])):
        db_session.add(User(user_id=identifier, username=identifier, display_name=identifier,
                            role="store", store_scope_mode="specified", status="active",
                            is_initialized=True, password_hash=hash_password_pbkdf2("secret")))
        db_session.add_all(UserStoreScope(user_id=identifier, store_id=store_id) for store_id in ids)
    db_session.commit()
    _login(client, "db-admin")
    rows = client.get("/api/v1/admin/accounts").json()["data"]["rows"]
    assert [row["user_id"] for row in rows] == ["store-user"]
    assert [row["store_id"] for row in client.get("/api/v1/admin/account-store-catalog").json()["data"]["stores"]] == ["store-1"]
    for identifier in ("outside", "mixed", "empty", "db-admin"):
        assert client.put(f"/api/v1/admin/accounts/{identifier}", json=payload(identifier)).status_code == 404
        assert client.post(f"/api/v1/admin/accounts/{identifier}/reset-password",
                           json={"password": "changed", "password_confirm": "changed"}).status_code == 404
        assert client.put(f"/api/v1/admin/accounts/{identifier}/page-permissions",
                          json={"extra_allow": [], "extra_deny": []}).status_code == 404
    assert client.post("/api/v1/admin/accounts", json=payload()).status_code == 200
    # JSON null organization bindings must not make newly created accounts disappear.
    created = next(row for row in client.get("/api/v1/admin/accounts").json()["data"]["rows"] if row["username"] == "created")
    assert client.post(f"/api/v1/admin/accounts/{created['user_id']}/reset-password",
                       json={"password": "changed", "password_confirm": "changed"}).status_code == 200
    assert client.post("/api/v1/admin/accounts", json=payload("outside-create", store_ids=["store-2"])).status_code == 403
    assert client.post("/api/v1/admin/accounts", json=payload("admin-create", role="admin", store_scope_mode="all", store_ids=[])).status_code == 403
    assert client.put("/api/v1/admin/accounts/store-user", json=payload("store-user", store_ids=["store-2"])).status_code == 403
    assert db_session.get(User, "store-user").store_scope_mode == "specified"
    assert client.delete("/api/v1/admin/accounts/store-user").status_code == 405


def test_scoped_bulk_template_preview_and_commit_are_scope_safe(client, db_session):
    scoped_admin(db_session)
    _login(client, "db-admin")
    template = client.get("/api/v1/admin/account-bulk-import/template")
    assert template.status_code == 200
    book = load_workbook(BytesIO(template.content))
    assert next(iter(book["账号开通"].data_validations.dataValidation)).formula1 == '"门店账号"'
    assert [row[0] for row in book["可选门店"].iter_rows(min_row=2, values_only=True)] == ["store-1"]
    content = workbook([["ok", "门店", "门店账号", "", "", "", "", "store-1"],
                        ["outside", "越权", "门店账号", "", "", "", "", "store-2"],
                        ["manager", "管理员", "管理员"]])
    files = {"file": ("accounts.xlsx", content)}
    preview = client.post("/api/v1/admin/account-bulk-import/preview", files=files).json()["data"]
    assert preview["allowed_account_types"] == ["门店账号"]
    assert len(preview["errors"]) == 2
    before = db_session.scalar(select(func.count()).select_from(User))
    assert client.post("/api/v1/admin/account-bulk-import/commit", files=files,
                       data={"digest": preview["digest"]}).status_code == 422
    assert db_session.scalar(select(func.count()).select_from(User)) == before
    content = workbook([["bulk-ok", "门店", "门店账号", "", "", "", "", "store-1"]])
    files = {"file": ("accounts.xlsx", content)}
    preview = client.post("/api/v1/admin/account-bulk-import/preview", files=files).json()["data"]
    assert not preview["errors"]
    result = client.post("/api/v1/admin/account-bulk-import/commit", files=files, data={"digest": preview["digest"]})
    assert result.status_code == 200
    user = db_session.scalar(select(User).where(User.username == "bulk-ok"))
    assert user.role == "store" and user.store_scope_mode == "specified"
    assert db_session.scalars(select(UserStoreScope.store_id).where(UserStoreScope.user_id == user.user_id)).all() == ["store-1"]


@pytest.mark.parametrize("path,method", [
    ("/admin/sku-fee-rules", "post"), ("/admin/sku-fee-rules/rebuild", "post"),
    ("/admin/sku-fee-rule-imports", "post"), ("/admin/sku-fee-rule-imports/missing/commit", "post"),
    ("/admin/settlement-scope-rules", "post"), ("/admin/sku-products/sku", "put"),
    ("/admin/sku-product-imports", "post"), ("/admin/sku-rules", "put"),
    ("/admin/non-commission-owner-accounts", "put"), ("/admin/product-type-visibility", "put"),
    ("/admin/sync/run", "post"), ("/admin/sync/config", "put"), ("/admin/product-sync-runs", "post"),
])
def test_scoped_admin_cannot_write_global_resources(client, db_session, path, method):
    scoped_admin(db_session)
    _login(client, "db-admin")
    response = getattr(client, method)("/api/v1" + path, json={})
    assert response.status_code == 403
    assert response.json()["detail"]["code"] == "GLOBAL_WRITE_REQUIRED"
    assert db_session.scalar(select(func.count()).select_from(JobRun)) == 0


def test_fixed_admin_defaults_and_last_global_admin_are_protected(client):
    _login(client, "system-admin", "test-password")
    for changes in ({"status": "disabled"}, {"role": "store", "store_scope_mode": "specified", "store_ids": ["store-1"]},
                    {"role": "admin", "store_scope_mode": "specified", "store_ids": ["store-1"]}):
        data = payload("db-admin", role="admin", store_scope_mode="all", store_ids=[], **{})
        data.update(changes)
        response = client.put("/api/v1/admin/accounts/db-admin", json=data)
        assert response.status_code == 409
        assert response.json()["detail"]["code"] == "LAST_GLOBAL_ADMIN_REQUIRED"
    response = client.put("/api/v1/admin/access-control/roles/admin", json={"page_keys": [], "confirmed": True})
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "ADMIN_PERMISSIONS_FIXED"
    assert client.post("/api/v1/admin/accounts", json=payload("backup", role="admin", store_scope_mode="all", store_ids=[])).status_code == 200
    assert client.put("/api/v1/admin/accounts/db-admin", json=payload("db-admin", role="admin", store_scope_mode="all", store_ids=[], status="disabled")).status_code == 200


def test_account_changes_invalidate_web_and_cli_generation(client, db_session):
    _login(client, "store-user")
    previous_cookie = client.cookies.get("dy_session")
    user = db_session.get(User, "store-user")
    versions = user.auth_version, user.auth_generation
    _login(client, "system-admin", "test-password")
    response = client.post("/api/v1/admin/accounts/store-user/reset-password", json={"password": "new-secret", "password_confirm": "new-secret"})
    assert response.status_code == 200
    db_session.refresh(user)
    assert user.auth_version > versions[0] and user.auth_generation > versions[1]
    client.cookies.set("dy_session", previous_cookie)
    assert client.get("/api/v1/auth/me").status_code == 401
    audit = db_session.scalars(select(AccountPermissionAuditLog).where(AccountPermissionAuditLog.action == "account.password_reset")).one()
    assert "new-secret" not in repr(audit.after_json)
