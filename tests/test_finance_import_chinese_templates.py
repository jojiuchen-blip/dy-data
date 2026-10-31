"""Chinese files exercise the real upload, preview, commit and fresh-read path."""

from csv import DictReader, writer
from datetime import datetime, timezone
from io import BytesIO, StringIO
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from apps.api.dy_api.finance_import_templates import (
    FINANCE_IMPORT_EXAMPLE_MARKER,
    FINANCE_IMPORT_FIELD_LABELS,
    FINANCE_IMPORT_TEMPLATE_FIELDS,
    finance_import_headers,
)
from apps.api.dy_api.models import FinanceImportBatch, InvoiceRecord, PromotionInvoice, StoreFinanceProfile
from test_api_finance_imports import _commit, _login, _upload, client  # noqa: F401
from test_dydata104_finance_state_matrix import CONTENTS, _seed, db_session  # noqa: F401


def _post_file(client, kind, content, suffix="xlsx"):
    return client.post(
        "/api/v1/admin/finance-imports",
        data={"importType": kind, "statementMonth": "2026-08"},
        files={"file": (f"input.{suffix}", content)},
        headers={"Idempotency-Key": uuid4().hex},
    )


def _save(workbook):
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


@pytest.mark.parametrize("kind", FINANCE_IMPORT_TEMPLATE_FIELDS)
def test_chinese_workbook_has_blank_data_and_one_isolated_text_example(client, db_session, kind):
    _login(client)
    response = client.get(f"/api/v1/admin/finance-imports/templates/{kind}")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert response.headers["content-disposition"].endswith(".xlsx")
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["数据导入", "填写示例", "填写说明"]
    assert list(workbook["数据导入"].values) == [tuple(finance_import_headers(kind))]
    assert workbook["数据导入"].column_dimensions["A"].number_format == "@"
    example_rows = list(workbook["填写示例"].values)
    assert len(example_rows) == 2
    assert example_rows[0] == (*finance_import_headers(kind), FINANCE_IMPORT_EXAMPLE_MARKER)
    assert "虚构示例" in example_rows[1][-1]
    assert all(cell.data_type == "s" for cell in workbook["填写示例"][2] if cell.value)
    assert any("分" in str(row) for row in workbook["填写说明"].values) if "AmountCent" in str(FINANCE_IMPORT_TEMPLATE_FIELDS[kind]) else True
    # Unmodified template, even with the example sheet active, writes no batch or target.
    workbook.active = 1
    uploaded = _post_file(client, kind, _save(workbook))
    assert uploaded.status_code == 422
    assert uploaded.json()["detail"]["code"] == "IMPORT_EMPTY_FILE"
    assert db_session.scalar(select(func.count()).select_from(FinanceImportBatch)) == 0
    assert db_session.scalar(select(func.count()).select_from(StoreFinanceProfile)) == 0


@pytest.mark.parametrize("kind", FINANCE_IMPORT_TEMPLATE_FIELDS)
@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_chinese_files_commit_each_formal_family_and_reopen_stable_values(client, db_session, kind, file_format):
    _login(client)
    _seed(db_session)
    english = next(DictReader(StringIO(CONTENTS[kind])))
    values = [english[field] for field in FINANCE_IMPORT_TEMPLATE_FIELDS[kind]]
    values = [{"APPROVED_SETTLED": "审核通过并已结算", "CONFIRMED": "已确认"}.get(value, value) for value in values]
    if file_format == "xlsx":
        response = client.get(f"/api/v1/admin/finance-imports/templates/{kind}")
        workbook = load_workbook(BytesIO(response.content))
        workbook["数据导入"].append(values)
        workbook.active = 1  # A saved active example must never become the data source.
        content = _save(workbook)
    else:
        buffer = StringIO()
        csv = writer(buffer)
        csv.writerow(finance_import_headers(kind))
        csv.writerow(values)
        content = buffer.getvalue().encode("utf-8-sig")
    uploaded = _post_file(client, kind, content, file_format)
    assert uploaded.status_code == 200, uploaded.text
    batch = uploaded.json()["data"]
    assert batch["totalRows"] == 1 and batch["successRows"] == 1 and batch["errorRows"] == 0
    assert batch["scenario"] == "FIRST_IMPORT_READY"
    preview = client.get(f"/api/v1/admin/finance-imports/{batch['batchId']}/rows").json()["data"]
    assert preview["total"] == 1 and preview["list"][0]["preview"]["stateTransition"]["allowed"]
    summary = preview["previewSummary"]
    assert sum(summary[key] for key in ("createRowCount", "updateRowCount", "unchangedRowCount", "failedRowCount", "unavailableRowCount")) == 1
    assert summary["writtenRowCount"] == 0 and summary["pendingRowCount"] == 1
    assert _commit(client, batch, uuid4().hex).status_code == 200
    with Session(db_session.get_bind()) as fresh:
        if kind in {"BASIC_INFO", "SAP_CONFIRMATION"}:
            target = fresh.scalar(select(StoreFinanceProfile).where(StoreFinanceProfile.store_id == "store-1", StoreFinanceProfile.is_current.is_(True)))
            assert target.sap_code == (english["sapCode"] if kind == "BASIC_INFO" else english["finalSapCode"])
            if kind == "SAP_CONFIRMATION":
                assert target.factory_confirmed is True
        elif kind == "PROMOTION_FACTORY_RESULT":
            target = fresh.scalar(select(PromotionInvoice).where(PromotionInvoice.is_current.is_(True)))
            assert target.invoice_amount_cent == 1200 and target.invoice_status == 3
        else:
            target = fresh.scalar(select(InvoiceRecord).where(InvoiceRecord.is_current.is_(True)))
            assert target.factory_deduction_amount_cent == 2500 and target.invoice_status == 3
    committed_summary = client.get(f"/api/v1/admin/finance-imports/{batch['batchId']}/rows").json()["data"]["previewSummary"]
    assert committed_summary["writtenRowCount"] == 1 and committed_summary["pendingRowCount"] == 0
    # English old-template replay still resolves to the same canonical request and no new write.
    repeated = _upload(client, kind, CONTENTS[kind])
    assert repeated.json()["data"]["scenario"] == "NO_CHANGE"
    repeated_id = repeated.json()["data"]["batchId"]
    no_change = client.get(f"/api/v1/admin/finance-imports/{repeated_id}/rows").json()["data"]["previewSummary"]
    assert no_change["unchangedRowCount"] == 1 and no_change["writtenRowCount"] == 0


@pytest.mark.parametrize("headers", [
    "门店ID,storeId,门店名称,SAP编码,资料导入时间（含时区）",
    "storeId,storeId,storeName,sapCode,importedAt",
    "门店ID,门店名称,错误SAP列,资料导入时间（含时区）",
])
def test_duplicate_alias_and_unknown_header_cannot_silently_overwrite(client, db_session, headers):
    _login(client)
    response = _upload(client, "BASIC_INFO", headers + "\nstore-1,store-2,Store One,SAP-1,2026-08-21T10:00:00+08:00\n")
    assert response.status_code == 422
    assert db_session.scalar(select(func.count()).select_from(StoreFinanceProfile)) == 0
    assert db_session.scalar(select(func.count()).select_from(FinanceImportBatch)) == 0


def test_example_only_or_ambiguous_worksheets_are_rejected(client, db_session):
    _login(client)
    response = client.get("/api/v1/admin/finance-imports/templates/BASIC_INFO")
    workbook = load_workbook(BytesIO(response.content))
    workbook.remove(workbook["数据导入"])
    assert _post_file(client, "BASIC_INFO", _save(workbook)).status_code == 422
    workbook["填写示例"].title = "数据导入"
    assert _post_file(client, "BASIC_INFO", _save(workbook)).status_code == 422
    workbook = Workbook()
    workbook.active.append(["storeId", "storeName", "sapCode", "importedAt"])
    workbook.create_sheet("另一数据表").append(["storeId", "storeName"])
    assert _post_file(client, "BASIC_INFO", _save(workbook)).status_code == 422
    assert db_session.scalar(select(func.count()).select_from(StoreFinanceProfile)) == 0


def test_long_numeric_excel_identifier_is_blocked_but_text_keeps_leading_zero(client):
    _login(client)
    workbook = Workbook()
    workbook.active.append(finance_import_headers("BASIC_INFO"))
    workbook.active.append(["store-1", "Store One", 12345678901234567890, "2026-08-21T10:00:00+08:00"])
    uploaded = _post_file(client, "BASIC_INFO", _save(workbook))
    assert uploaded.status_code == 200
    batch = uploaded.json()["data"]
    assert batch["scenario"] == "BATCH_VALIDATION_FAILED"
    errors = client.get(f"/api/v1/admin/finance-imports/{batch['batchId']}").json()["data"]["errors"]["list"]
    assert any("精度" in error["reason"] for error in errors)
    workbook.active.cell(2, 3, "000012345678901234567890")
    good = _post_file(client, "BASIC_INFO", _save(workbook))
    assert good.json()["data"]["scenario"] == "FIRST_IMPORT_READY"


@pytest.mark.parametrize("sap", ["SAP-FINAL", ""])
def test_basic_metadata_import_preserves_existing_sap_confirmation_and_refreshes_formal_queries(client, db_session, sap):
    _login(client)
    timestamp = datetime(2026, 8, 21, 3, tzinfo=timezone.utc)
    db_session.add(StoreFinanceProfile(
        profile_id="confirmed-old", store_id="store-1", profile_type=1, source_type=1,
        version_no=1, is_current=True, store_name_snapshot="Store One", sap_code="SAP-FINAL",
        initial_sap_code="SAP-OLD", service_store_code="0000123", factory_confirmed=True, confirmed_at=timestamp,
    ))
    db_session.commit()
    uploaded = _upload(client, "BASIC_INFO",
        "门店ID,门店名称,SAP编码,资料导入时间（含时区）,SAP编码确认状态\n"
        f"store-1,Store One,{sap},2026-08-22T10:00:00+08:00,REJECTED\n")
    batch = uploaded.json()["data"]
    assert batch["scenario"] == "FIRST_IMPORT_READY"
    assert _commit(client, batch, uuid4().hex).status_code == 200
    with Session(db_session.get_bind()) as fresh:
        current = fresh.scalar(select(StoreFinanceProfile).where(StoreFinanceProfile.is_current.is_(True)))
        old = fresh.scalar(select(StoreFinanceProfile).where(StoreFinanceProfile.profile_id == "confirmed-old"))
        assert not old.is_current and old.factory_confirmed is True
        assert current.sap_code == "SAP-FINAL" and current.factory_confirmed is True
        assert current.confirmed_at.replace(tzinfo=timezone.utc) == timestamp
        assert current.initial_sap_code == "SAP-OLD" and current.service_store_code == "0000123"
    query = {"month": "2026-08", "feeDirection": "PROMOTION", "metricScope": "MONTH", "q": "SAP-FINAL"}
    items = client.get("/api/v1/admin/finance/stores", params=query).json()["data"]["list"]
    assert len(items) == 1 and items[0]["effectiveSapCode"] == "SAP-FINAL" and items[0]["sapStatus"] == "CONFIRMED"
    exported = client.get("/api/v1/admin/finance/stores/export", params=query)
    rows = list(DictReader(StringIO(exported.content.decode("utf-8-sig"))))
    assert len(rows) == 1 and rows[0]["有效SAP编码"] == "SAP-FINAL"
    assert "sap_status" not in rows[0] and "SAP确认状态" not in rows[0]


def test_new_sap_value_keeps_old_confirmation_only_in_history_and_preview_explains(client, db_session):
    _login(client)
    db_session.add(StoreFinanceProfile(profile_id="old", store_id="store-1", profile_type=1,
        source_type=1, version_no=1, is_current=True, store_name_snapshot="Store One", sap_code="OLD",
        factory_confirmed=True, confirmed_at=datetime(2026, 8, 21, tzinfo=timezone.utc)))
    db_session.commit()
    uploaded = _upload(client, "BASIC_INFO", "storeId,storeName,sapCode,importedAt\nstore-1,Store One,NEW,2026-08-22T10:00:00+08:00\n")
    batch = uploaded.json()["data"]
    preview = client.get(f"/api/v1/admin/finance-imports/{batch['batchId']}/rows").json()["data"]["list"][0]["preview"]
    assert any("不自动获得厂家确认" in text for text in preview["allowedChanges"])
    assert _commit(client, batch, uuid4().hex).status_code == 200
    with Session(db_session.get_bind()) as fresh:
        current = fresh.scalar(select(StoreFinanceProfile).where(StoreFinanceProfile.is_current.is_(True)))
        old = fresh.scalar(select(StoreFinanceProfile).where(StoreFinanceProfile.profile_id == "old"))
        assert current.sap_code == "NEW" and current.factory_confirmed is None and current.confirmed_at is None
        assert old.factory_confirmed is True and old.confirmed_at is not None


def test_formula_or_later_malformed_row_rolls_back_the_whole_upload(client, db_session):
    _login(client)
    workbook = Workbook()
    workbook.active.append(finance_import_headers("BASIC_INFO"))
    workbook.active.append(["store-1", "Store One", "SAP-1", "2026-08-21T10:00:00+08:00"])
    workbook.active.append(["store-2", "Store Two", '=CONCAT("SAP", "-2")', "2026-08-21T10:00:00+08:00"])
    response = _post_file(client, "BASIC_INFO", _save(workbook))
    assert response.status_code == 422
    assert "公式" in response.json()["detail"]["message"]
    assert db_session.scalar(select(func.count()).select_from(FinanceImportBatch)) == 0
    assert db_session.scalar(select(func.count()).select_from(StoreFinanceProfile)) == 0
    content = "storeId,storeName,sapCode,importedAt\nstore-1,Store One,SAP-1,2026-08-21T10:00:00+08:00\nstore-2,Store Two,SAP-2,2026-08-21T10:00:00+08:00,unexpected\n"
    assert _upload(client, "BASIC_INFO", content).status_code == 422
    assert db_session.scalar(select(func.count()).select_from(FinanceImportBatch)) == 0


def test_missing_basic_sap_column_preserves_value_without_changing_confirmation(client, db_session):
    _login(client)
    db_session.add(StoreFinanceProfile(profile_id="preserve", store_id="store-1", profile_type=1,
        source_type=1, version_no=1, is_current=True, store_name_snapshot="Store One", sap_code="0000123",
        factory_confirmed=False, confirmed_at=datetime(2026, 8, 21, tzinfo=timezone.utc)))
    db_session.commit()
    response = _upload(client, "BASIC_INFO", "门店ID,门店名称,资料导入时间（含时区）\nstore-1,Store One,2026-08-22T10:00:00+08:00\n")
    batch = response.json()["data"]
    assert _commit(client, batch, uuid4().hex).status_code == 200
    with Session(db_session.get_bind()) as fresh:
        current = fresh.scalar(select(StoreFinanceProfile).where(StoreFinanceProfile.is_current.is_(True)))
        assert current.sap_code == "0000123" and current.factory_confirmed is False
