from test_api_access_control import client, _login
from apps.api.dy_api.models import User, AccountPermissionAuditLog
from apps.api.dy_api.account_bulk_import import account_template
from openpyxl import load_workbook
from io import BytesIO
from sqlalchemy import select
import json


def workbook(rows):
    book = load_workbook(BytesIO(account_template()))
    for number, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            book['账号开通'].cell(number, column, value)
    output = BytesIO(); book.save(output)
    return output.getvalue()


def test_batch_preview_create_passwords_and_repeat_protection(client, db_session):
    _login(client, 'system-admin', 'test-password')
    content = workbook([
        ['00123', '张店长', '门店账号', '', '', '', '', 'store-1;store-2'],
        ['manager_li', '李经理', '管理员'],
    ])
    files = {'file': ('accounts.xlsx', content, 'application/octet-stream')}
    preview = client.post('/api/v1/admin/account-bulk-import/preview', files=files)
    assert preview.status_code == 200, preview.text
    data = preview.json()['data']
    assert not data['errors'] and len(data['rows']) == 2
    assert db_session.scalar(select(User).where(User.username == '00123')) is None
    created = client.post('/api/v1/admin/account-bulk-import/commit', files=files, data={'digest': data['digest']})
    assert created.status_code == 200, created.text if created.status_code != 200 else ''
    assert created.headers['cache-control'] == 'no-store'
    book = load_workbook(BytesIO(created.content))
    credentials = list(book.active.iter_rows(min_row=2, values_only=True))
    assert credentials[0][0] == '00123'
    assert len(credentials[0][3]) >= 16 and credentials[0][3] != credentials[1][3]
    audits = db_session.scalars(select(AccountPermissionAuditLog)).all()
    for audit in audits:
        assert credentials[0][3] not in repr(audit.__dict__)
    repeat = client.post('/api/v1/admin/account-bulk-import/commit', files=files, data={'digest': data['digest']})
    assert repeat.status_code == 422
    _login(client, '00123', credentials[0][3])
    assert set(client.get('/api/v1/auth/me').json()['data']['store_ids']) == {'store-1', 'store-2'}
    assert client.get('/api/v1/admin/account-bulk-import/template').status_code == 403


def test_invalid_batch_is_atomic_and_changed_file_refused(client, db_session):
    _login(client, 'system-admin', 'test-password')
    content = workbook([['valid_one', '有效', '门店账号', '', '', '', '', 'store-1'], ['bad_one', '无效', '门店账号', '', '', '', '', 'missing']])
    files = {'file': ('accounts.xlsx', content, 'application/octet-stream')}
    preview = client.post('/api/v1/admin/account-bulk-import/preview', files=files).json()['data']
    assert preview['errors'][0]['row'] == 3
    assert client.post('/api/v1/admin/account-bulk-import/commit', files=files, data={'digest': preview['digest']}).status_code == 422
    assert db_session.scalar(select(User).where(User.username == 'valid_one')) is None
    assert client.post('/api/v1/admin/account-bulk-import/commit', files=files, data={'digest': 'changed'}).status_code == 409


def test_bulk_rejects_ambiguous_identifiers_and_scope_columns(client, db_session):
    _login(client, 'system-admin', 'test-password')
    db_session.get(User, 'store-user').external_account_id = 'existing-alias'
    db_session.commit()
    content = workbook([
        ['a  b', '一', '门店账号', '', '', '', '', 'store-1'],
        ['a b', '二', '门店账号', '', '', '', '', 'store-2'],
        ['existing-alias', '三', '管理员'],
        ['wide', '四', '管理员', '某集团'],
    ])
    preview = client.post('/api/v1/admin/account-bulk-import/preview', files={'file': ('accounts.xlsx', content)}).json()['data']
    assert preview['rows'][0]['username'] == 'a b'
    assert [row['row'] for row in preview['errors']] == [3, 4, 5]
